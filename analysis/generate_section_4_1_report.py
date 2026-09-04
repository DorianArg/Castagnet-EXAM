"""Génère le rapport Word final du §4.1 CastagNet à partir des artefacts validés."""

from __future__ import annotations

import json
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from PIL import Image, ImageOps, ImageDraw
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
OUTPUT = REPORTS / "Rapport_CastagNet_4_1_Qualite_Dataset.docx"
ASSETS = ROOT / "analysis" / "report_assets"
OUT = ROOT / "analysis" / "output"
VIDEO = ROOT / "analysis" / "extracted_video"

NAVY = "17324D"
TEAL = "087E8B"
GOLD = "D9A441"
PALE = "EAF3F4"
LIGHT = "F4F6F8"
MID = "D6DEE5"
RED = "A33A3A"
GREEN = "317256"
INK = "24313D"
WHITE = "FFFFFF"

LABEL_COLORS = {
    "Vide": "#8A98A6",
    "Conforme": "#2A9D8F",
    "PIETRA": "#E9C46A",
    "NON Conforme": "#E76F51",
}


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = MID, size: str = "4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_table(doc: Document, headers: list[str], rows: list[list[object]], widths=None,
              font_size: float = 8.2) -> object:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, value in enumerate(headers):
        cell = hdr.cells[idx]
        shade(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(value))
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(WHITE)
        r.font.size = Pt(font_size)
    for ridx, row_values in enumerate(rows):
        row = table.add_row()
        for cidx, value in enumerate(row_values):
            cell = row.cells[cidx]
            if ridx % 2:
                shade(cell, LIGHT)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if cidx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            r.font.size = Pt(font_size)
        if widths:
            for cidx, width in enumerate(widths):
                row.cells[cidx].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text)


def add_picture(doc: Document, path: Path, caption: str, width_cm: float = 15.5) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    add_caption(doc, caption)


def add_lead(doc: Document, label: str, text: str, color: str = TEAL) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(label + " — ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(color)
    p.add_run(text)


def add_callout(doc: Document, title: str, text: str, fill: str = PALE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_border(cell, TEAL, "8")
    set_cell_margins(cell, top=130, start=170, bottom=130, end=170)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title + "\n")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def bullet(doc: Document, text: str, level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(text, style=style)
    p.paragraph_format.space_after = Pt(2)


def heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.75)
    section.bottom_margin = Cm(1.65)
    section.left_margin = Cm(1.85)
    section.right_margin = Cm(1.75)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.6)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True

    for name, size, color in (("Title", 27, NAVY), ("Subtitle", 13, TEAL)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)

    for level, size in ((1, 17), (2, 13), (3, 10.5)):
        style = styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else TEAL)
        style.paragraph_format.space_before = Pt(10 if level > 1 else 0)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
        if level == 1:
            style.paragraph_format.page_break_before = True

    caption = styles["Caption"]
    caption.font.name = "Aptos"
    caption.font.size = Pt(8)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string("566573")
    caption.paragraph_format.space_after = Pt(7)
    caption.paragraph_format.keep_with_next = False

    if "Small note" not in styles:
        note = styles.add_style("Small note", WD_STYLE_TYPE.PARAGRAPH)
        note.font.name = "Aptos"
        note.font.size = Pt(8)
        note.font.color.rgb = RGBColor.from_string("566573")
        note.paragraph_format.space_after = Pt(4)

    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = "CastagNet — Qualité et complétude de la donnée"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.runs[0].font.size = Pt(8)
        header.runs[0].font.color.rgb = RGBColor.from_string("6B7782")
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("MSc BIHAR  |  §4.1  |  ")
        add_field(footer, "PAGE", "1")

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    lang = OxmlElement("w:themeFontLang")
    lang.set(qn("w:val"), "fr-FR")
    settings.append(lang)
    return doc


def configure_plotting() -> None:
    plt.rcParams.update({
        "font.family": "DejaVu Sans",
        "font.size": 9,
        "axes.titlesize": 11,
        "axes.titleweight": "bold",
        "axes.labelcolor": "#24313D",
        "text.color": "#24313D",
        "axes.edgecolor": "#B4C0CA",
        "figure.facecolor": "white",
    })


def save_figure(fig, name: str) -> Path:
    ASSETS.mkdir(parents=True, exist_ok=True)
    path = ASSETS / name
    fig.savefig(path, dpi=210, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def chart_workflow() -> Path:
    fig, ax = plt.subplots(figsize=(11.5, 4.2))
    ax.axis("off")
    stages = [
        ("Audit", "Fichiers, labels,\nartefacts"),
        ("Vidéo", "Extraction et\ncrops manuels"),
        ("Relations T/B", "Temporalité vidéo\nnamespace historique"),
        ("Décision", "Stratégies et\nrevue ciblée"),
        ("Manifeste", "effective_label\nsplit_group_id"),
    ]
    xs = np.linspace(0.08, 0.92, len(stages))
    for i, ((title, subtitle), x) in enumerate(zip(stages, xs)):
        color = "#087E8B" if i not in (2, 3) else "#D9A441"
        circle = plt.Circle((x, 0.62), 0.075, color=color, transform=ax.transAxes)
        ax.add_patch(circle)
        ax.text(x, 0.62, str(i + 1), ha="center", va="center", color="white",
                weight="bold", fontsize=13, transform=ax.transAxes)
        ax.text(x, 0.39, title, ha="center", weight="bold", fontsize=10,
                transform=ax.transAxes)
        ax.text(x, 0.22, subtitle, ha="center", va="center", fontsize=8.5,
                transform=ax.transAxes)
        if i < len(stages) - 1:
            ax.annotate("", xy=(xs[i + 1] - 0.085, 0.62), xytext=(x + 0.085, 0.62),
                        arrowprops=dict(arrowstyle="->", color="#60717F", lw=1.6),
                        xycoords=ax.transAxes)
    ax.text(0.5, 0.91, "Du diagnostic de qualité à une référence ML traçable",
            ha="center", fontsize=13, weight="bold", transform=ax.transAxes)
    return save_figure(fig, "workflow_4_1.png")


def chart_label_referentials() -> Path:
    labels = ["Conforme", "NON Conforme", "PIETRA", "Vide"]
    historical = [14818, 9169, 11267, 0]
    principal = [10836, 5659, 6033, 12726]
    x = np.arange(len(labels))
    fig, ax = plt.subplots(figsize=(9.2, 4.8))
    w = 0.36
    b1 = ax.bar(x - w / 2, historical, w, label="label_filename (historique)", color="#8A98A6")
    b2 = ax.bar(x + w / 2, principal, w, label="label_principal (cible actuelle)", color="#087E8B")
    ax.bar_label(b1, padding=2, fontsize=8)
    ax.bar_label(b2, padding=2, fontsize=8)
    ax.set_xticks(x, labels)
    ax.set_ylabel("Nombre d’images")
    ax.set_title("Deux référentiels de labels à ne pas confondre")
    ax.legend(frameon=False, ncols=2, loc="upper center")
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="y", alpha=.15)
    return save_figure(fig, "label_referentials.png")


def chart_annual() -> Path:
    labels = ["Conforme", "NON Conforme", "PIETRA", "Vide"]
    values = {"2025": [7606, 3594, 2596, 6470], "2026": [3230, 2065, 3437, 6256]}
    fig, axes = plt.subplots(1, 2, figsize=(9.5, 4.3), sharey=True)
    for ax, (year, vals) in zip(axes, values.items()):
        bars = ax.bar(labels, vals, color=[LABEL_COLORS[x] for x in labels])
        ax.bar_label(bars, padding=2, fontsize=8)
        ax.set_title(f"{year} — n={sum(vals):,}".replace(",", " "))
        ax.tick_params(axis="x", rotation=20)
        ax.spines[["top", "right"]].set_visible(False)
        ax.grid(axis="y", alpha=.15)
    axes[0].set_ylabel("Nombre d’images")
    fig.suptitle("La composition des classes varie fortement entre 2025 et 2026", weight="bold")
    return save_figure(fig, "annual_distribution.png")


def chart_acquisition() -> Path:
    cameras = np.arange(1, 7)
    counts = [5811, 5953, 6157, 5792, 5913, 5628]
    fig, axes = plt.subplots(1, 2, figsize=(9.5, 4.2), gridspec_kw={"width_ratios": [1.65, 1]})
    bars = axes[0].bar(cameras, counts, color="#087E8B")
    axes[0].bar_label(bars, padding=2, fontsize=8)
    axes[0].set_ylim(0, 6800)
    axes[0].set_xticks(cameras)
    axes[0].set_xlabel("Emplacement caméra")
    axes[0].set_ylabel("Images")
    axes[0].set_title("Répartition par caméra")
    axes[0].spines[["top", "right"]].set_visible(False)
    axes[1].pie([17065, 18189], labels=["Top\n17 065", "Bottom\n18 189"],
                colors=["#D9A441", "#17324D"], autopct="%1.2f %%", startangle=90,
                textprops={"fontsize": 8.5, "color": "#24313D"})
    axes[1].set_title("Positions")
    return save_figure(fig, "acquisition_balance.png")


def chart_video_pairing() -> Path:
    fig, axes = plt.subplots(1, 2, figsize=(9.5, 4.2))
    statuses = ["Appariées", "Top seul", "Bottom seul", "Ambigu"]
    vals = [22, 2, 3, 1]
    colors = ["#2A9D8F", "#8A98A6", "#8A98A6", "#E9C46A"]
    bars = axes[0].bar(statuses, vals, color=colors)
    axes[0].bar_label(bars, padding=2)
    axes[0].tick_params(axis="x", rotation=20)
    axes[0].set_title("Sortie de l’alignement")
    axes[0].set_ylabel("Nombre")
    axes[0].spines[["top", "right"]].set_visible(False)
    bars = axes[1].bar(["+6", "+7", "+8"], [5, 17, 0], color=["#8A98A6", "#087E8B", "#D9A441"])
    axes[1].bar_label(bars, padding=2)
    axes[1].set_title("Décalage des 22 paires certaines")
    axes[1].set_xlabel("Delta en frames")
    axes[1].spines[["top", "right"]].set_visible(False)
    fig.suptitle("Appariement vidéo : stabilité autour de +7 frames", weight="bold")
    return save_figure(fig, "video_pairing.png")


def make_ambiguous_montage() -> Path:
    paths = [
        VIDEO / "top" / "crops" / "T_frame_000572_crop_01.jpg",
        VIDEO / "top" / "crops" / "T_frame_000572_crop_02.jpg",
        VIDEO / "bottom" / "crops" / "B_frame_000579_crop_01.jpg",
    ]
    labels = ["Top — candidat 1", "Top — candidat 2", "Bottom — unique"]
    tile_w, tile_h = 330, 300
    canvas = Image.new("RGB", (tile_w * 3, tile_h + 55), "white")
    draw = ImageDraw.Draw(canvas)
    for i, (path, label) in enumerate(zip(paths, labels)):
        img = Image.open(path).convert("RGB")
        img = ImageOps.contain(img, (tile_w - 30, tile_h - 30))
        x = i * tile_w + (tile_w - img.width) // 2
        y = 15 + (tile_h - img.height) // 2
        canvas.paste(img, (x, y))
        draw.text((i * tile_w + 18, tile_h + 12), label, fill="#24313D")
    ASSETS.mkdir(parents=True, exist_ok=True)
    out = ASSETS / "ambiguous_video_case.png"
    canvas.save(out, quality=95)
    return out


def chart_historical_keys() -> Path:
    keys = ["A", "B", "C", "D"]
    pairs = [5561, 16176, 12318, 932]
    coverage = [31.55, 91.77, 69.88, 5.29]
    fig, ax = plt.subplots(figsize=(8.7, 4.6))
    bars = ax.bar(keys, coverage, color=["#8A98A6", "#087E8B", "#D9A441", "#B9C1C8"])
    for bar, pct, count in zip(bars, coverage, pairs):
        ax.text(bar.get_x() + bar.get_width()/2, bar.get_height()+2,
                f"{pct:.2f} %\n{count:,} paires".replace(",", " "), ha="center", fontsize=8)
    ax.set_ylim(0, 108)
    ax.set_ylabel("Couverture des images (%)")
    ax.set_title("Comparaison des clés candidates d’appariement historique")
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="y", alpha=.15)
    return save_figure(fig, "historical_key_comparison.png")


def chart_strategies() -> Path:
    fig, axes = plt.subplots(1, 2, figsize=(10, 4.4))
    strategies = ["S0", "S1", "S2/S3a", "S3b"]
    counts = [35254, 35195, 29043, 22103]
    bars = axes[0].bar(strategies, counts, color=["#17324D", "#087E8B", "#D9A441", "#A33A3A"])
    axes[0].bar_label(bars, labels=[f"{x:,}".replace(",", " ") for x in counts], padding=2, fontsize=8)
    axes[0].set_ylim(0, 39000)
    axes[0].set_ylabel("Images conservées")
    axes[0].set_title("Volume selon la stratégie")
    axes[0].spines[["top", "right"]].set_visible(False)
    labels = ["Vide", "Conforme", "PIETRA"]
    x = np.arange(3)
    w = .36
    axes[1].bar(x-w/2, [36.10, 30.74, 17.11], w, label="S0", color="#17324D")
    axes[1].bar(x+w/2, [43.66, 27.28, 13.00], w, label="S2", color="#D9A441")
    axes[1].set_xticks(x, labels)
    axes[1].set_ylabel("Part de classe (%)")
    axes[1].set_title("Biais induit par S2")
    axes[1].legend(frameon=False)
    axes[1].spines[["top", "right"]].set_visible(False)
    fig.suptitle("Nettoyer davantage ne garantit pas une meilleure représentativité", weight="bold")
    return save_figure(fig, "strategy_comparison.png")


def chart_final() -> Path:
    fig, axes = plt.subplots(1, 2, figsize=(10, 4.4))
    labels = ["Vide", "Conforme", "PIETRA", "NON Conforme"]
    vals = [12724, 10832, 6034, 5664]
    axes[0].pie(vals, labels=labels, autopct="%1.2f %%", startangle=95,
                colors=[LABEL_COLORS[x] for x in labels], textprops={"fontsize": 8})
    axes[0].set_title("effective_label — 35 254 images")
    sources = ["Principal", "Confirmé humain", "Revue binaire", "Revue multiclasses"]
    source_vals = [35195, 43, 7, 9]
    bars = axes[1].barh(sources, source_vals, color=["#17324D", "#2A9D8F", "#D9A441", "#E76F51"])
    axes[1].set_xscale("symlog", linthresh=10)
    axes[1].bar_label(bars, padding=4, labels=[f"{x:,}".replace(",", " ") for x in source_vals])
    axes[1].set_title("Provenance du label final (échelle symlog)")
    axes[1].spines[["top", "right"]].set_visible(False)
    fig.suptitle("Dataset final : stabilité globale et corrections ciblées", weight="bold")
    return save_figure(fig, "final_dataset.png")


def generate_assets() -> dict[str, Path]:
    configure_plotting()
    return {
        "workflow": chart_workflow(),
        "labels": chart_label_referentials(),
        "annual": chart_annual(),
        "acquisition": chart_acquisition(),
        "video": chart_video_pairing(),
        "ambiguous": make_ambiguous_montage(),
        "keys": chart_historical_keys(),
        "strategies": chart_strategies(),
        "final": chart_final(),
    }


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CASTAGNET")
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor.from_string(TEAL)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Qualité et complétude\nde la donnée")
    p.paragraph_format.space_before = Pt(34)
    p.paragraph_format.space_after = Pt(12)

    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Rapport final — §4.1")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    r = p.add_run("Mise en Situation Professionnelle — MSc BIHAR\n")
    r.bold = True
    p.add_run("Compétences C28 et C31\n2 septembre 2026")

    doc.add_paragraph("\n")
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    values = [
        ("Périmètre", "Audit, extraction vidéo, correspondances T/B, arbitrage et manifeste final"),
        ("Dataset", "35 254 images historiques — 4 classes — 6 emplacements — vues Top/Bottom"),
        ("Décision finale", "100 % des images conservées ; 16 corrections humaines traçables"),
        ("Livrable ML", "analysis/output/final_training_manifest.csv"),
    ]
    for i, (left, right) in enumerate(values):
        shade(table.cell(i, 0), NAVY)
        shade(table.cell(i, 1), LIGHT if i % 2 else PALE)
        for cell in table.rows[i].cells:
            set_cell_border(cell)
            set_cell_margins(cell, 120, 150, 120, 150)
        p0 = table.cell(i, 0).paragraphs[0]
        r0 = p0.add_run(left)
        r0.bold = True
        r0.font.color.rgb = RGBColor.from_string(WHITE)
        table.cell(i, 1).paragraphs[0].add_run(right)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(40)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Document limité au §4.1 — aucun résultat de modélisation")
    r.italic = True
    r.font.color.rgb = RGBColor.from_string("6B7782")
    doc.add_page_break()


def add_executive_summary(doc: Document) -> None:
    p = doc.add_paragraph("Synthèse décisionnelle", style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph(
        "Le dataset est exploitable pour préparer l’apprentissage, à condition de conserver les distinctions "
        "entre labels historiques et cible actuelle, de traiter les associations Top/Bottom comme heuristiques "
        "et d’imposer des groupes communs lors des futurs splits. L’audit ne justifie pas une purge massive : "
        "les suppressions fondées sur chunk, multiple ou la provenance auraient déformé la distribution des classes."
    )
    add_table(doc, ["Question", "Réponse retenue"], [
        ["Dataset final", "35 254 images ; aucune exclusion"],
        ["Cible ML", "effective_label, dérivé sans modifier les CSV sources"],
        ["Corrections", "16 décisions humaines effectives sur 35 254 images"],
        ["Données difficiles", "multiple et chunk conservés, flags maintenus"],
        ["Top/Bottom vidéo", "Alignement temporel autour de +7 frames ; ambiguïtés non forcées"],
        ["Top/Bottom historique", "16 176 paires fortement soutenues mais heuristiques"],
        ["Protection future", "split_group_id commun aux deux vues d’une paire"],
        ["Limite centrale", "Absence d’identifiant physique de châtaigne et hétérogénéité historique"],
    ], widths=[4.5, 11.5], font_size=8.6)
    add_callout(doc, "Conclusion en une phrase",
                "Une donnée difficile mais comprise vaut mieux qu’un dataset artificiellement simplifié : "
                "le choix final conserve la réalité de production et rend les incertitudes explicites.")
    doc.add_page_break()


def add_toc(doc: Document) -> None:
    p = doc.add_paragraph("Sommaire", style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    toc = doc.add_paragraph()
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u', "Sommaire automatique — mise à jour à l’ouverture")
    doc.add_paragraph(
        "Le sommaire est un champ Word automatique fondé sur les niveaux de titres 1 à 3.",
        style="Small note",
    )
    doc.add_page_break()


def build_report() -> Document:
    # Chargements délibérés : ils garantissent que les sorties validées existent au moment du rapport.
    quality = load_json(OUT / "dataset_quality_summary.json")
    video_summary = load_json(VIDEO / "tb_pairing_summary.json")
    historical = load_json(OUT / "historical_tb_pairs_summary.json")
    adjudication = load_json(OUT / "masked_disagreements_adjudication_summary.json")
    final = load_json(OUT / "final_training_manifest_summary.json")
    assert quality["dataset_images"] == 35254
    assert video_summary["status_counts"]["matched_temporal"] == 22
    assert historical["pairs"] == 16176
    assert adjudication["multiclass_review_still_pending"] == 0
    assert final["total_images"] == 35254 and all(final["validations"].values())

    assets = generate_assets()
    doc = setup_document()
    props = doc.core_properties
    props.title = "Rapport CastagNet §4.1 — Qualité et complétude de la donnée"
    props.subject = "Audit du dataset, extraction vidéo, correspondances T/B et manifeste final"
    props.author = "Projet CastagNet — MSc BIHAR"
    props.keywords = "CastagNet, qualité des données, Top Bottom, labellisation, MSc BIHAR"

    add_cover(doc)
    add_executive_summary(doc)
    add_toc(doc)

    heading(doc, "1. Cadre, enjeu et méthode", 1)
    heading(doc, "1.1 Une donnée comprise avant tout entraînement", 2)
    doc.add_paragraph(
        "CastagNet s’inscrit dans une machine de tri de châtaignes sèches comportant six emplacements caméra, "
        "chacun observé par une vue du dessus (Top) et une vue du dessous (Bottom), soit douze flux simultanés. "
        "Les images sont classées en Conforme, NON Conforme, PIETRA ou Vide. La première exigence n’est donc "
        "pas de produire un modèle, mais de vérifier que la cible, l’unité statistique et les relations entre vues "
        "sont correctement comprises. Un score élevé sur une donnée mal structurée ne répondrait pas au besoin métier."
    )
    add_lead(doc, "Enjeu métier",
             "L’annexe GRPTMC privilégie la protection du lot Conforme : accepter un fruit NON Conforme ou PIETRA "
             "menace la qualité et la certification, tandis que rejeter un Conforme réduit surtout le rendement.")
    add_lead(doc, "Conséquence pour la suite",
             "Le §4.2 devra viser, sur Conforme, un rappel d’au moins 85 % et une précision d’au moins 95 %. "
             "Aucun résultat ML n’est présenté ici.")

    heading(doc, "1.2 Démarche d’analyse", 2)
    add_picture(doc, assets["workflow"],
                "Figure 1 — Workflow du §4.1, de l’audit au manifeste final.", 16.2)
    doc.add_paragraph(
        "La démarche combine deux niveaux : une observation directe de l’extrait vidéo, où la temporalité est "
        "disponible, et un audit structurel du dataset historique, où seul le nom de fichier subsiste. Dans les deux "
        "cas, une association n’est retenue que si son statut et ses limites peuvent être explicités."
    )

    heading(doc, "2. Diagnostic du dataset initial", 1)
    heading(doc, "2.1 Deux référentiels de labels", 2)
    add_lead(doc, "Constat",
             "Le nom de fichier contient un diagnostic historique, alors que labels_principal.csv contient la cible "
             "d’entraînement actuelle. Les deux champs ne répondent pas à la même question.")
    add_table(doc, ["Champ", "Rôle", "Usage retenu"], [
        ["label_filename", "Diagnostic historique encodé dans filename", "Audit et namespace historique uniquement"],
        ["label_principal", "Cible d’entraînement fournie", "Référence avant adjudication"],
        ["effective_label", "Cible finale dérivée", "Référence future du §4.2"],
    ], widths=[3.3, 6.2, 6.5])
    add_picture(doc, assets["labels"],
                "Figure 2 — Écart entre le diagnostic historique du filename et la cible actuelle.", 15.4)
    add_table(doc, ["Référentiel", "Conforme", "NON Conforme", "PIETRA", "Vide", "Total"], [
        ["label_filename", "14 818", "9 169", "11 267", "0", "35 254"],
        ["label_principal", "10 836", "5 659", "6 033", "12 726", "35 254"],
    ])
    add_lead(doc, "Interprétation",
             "Le rapport historique n’est pas nécessairement faux dans son propre référentiel ; il ne décrit simplement "
             "pas la cible ML actuelle. Utiliser ses chiffres comme vérité d’entraînement aurait notamment fait disparaître Vide.")
    add_lead(doc, "Décision", "Toutes les statistiques de cible reposent sur label_principal, puis sur effective_label après revue.")

    heading(doc, "2.2 Identité des images et artefact `_1` / `_2`", 2)
    doc.add_paragraph(
        "Les fichiers 2025 peuvent contenir un segment supplémentaire, par exemple "
        "2025_Conforme_1_Cam_B_1_17.jpg et 2025_Conforme_2_Cam_B_1_17.jpg. "
        "Cette information, présente sur 9 986 images, était reconnue mais non conservée dans les colonnes structurées."
    )
    add_lead(doc, "Choix",
             "filename reste l’identifiant canonique et unique ; batch_num est dérivé ; aucune déduplication n’est "
             "effectuée sur une combinaison partielle de métadonnées.")
    add_lead(doc, "Limite",
             "batch_num est un nom technique. En l’absence du générateur historique, aucune signification métier ne lui est attribuée.")

    heading(doc, "2.3 Provenance et statut de relecture", 2)
    add_table(doc, ["Indicateur", "Résultat", "Lecture correcte"], [
        ["Identifiants bruts non vides", "7", "Plusieurs variantes désignent le même annotateur"],
        ["Annotateurs canoniques", "4", "Regroupement statistique uniquement"],
        ["Nico", "15 518 (44,02 %)", "NicoG, Nico, nico et nico h"],
        ["reviewed=True", "35 254 (100 %)", "Valeur déclarée dans le CSV"],
        ["reviewed=True sans labeled_by", "8 033 (22,79 %)", "Provenance manquante, pas absence de relecture"],
    ])
    add_lead(doc, "Choix",
             "Les valeurs brutes de labeled_by sont conservées et une colonne labeled_by_canonical corrige seulement les statistiques.")
    add_callout(doc, "Contradiction non résolue",
                "Le sujet annonce environ 500 images 2026 restant à relire, mais les 35 254 lignes fournies ont reviewed=True. "
                "Aucun reliquat n’est inventé : le champ disponible ne permet pas d’identifier ces images. Un statut sans "
                "provenance complète limite toutefois sa valeur probante.", "FFF4E2")

    heading(doc, "2.4 Années, caméras et positions", 2)
    add_picture(doc, assets["annual"],
                "Figure 3 — Distribution de label_principal par année.", 15.8)
    add_table(doc, ["Année", "Conforme", "NON Conforme", "PIETRA", "Vide", "Total"], [
        ["2025", "7 606", "3 594", "2 596", "6 470", "20 266"],
        ["2026", "3 230", "2 065", "3 437", "6 256", "14 988"],
    ])
    add_lead(doc, "Constat",
             "La composition diffère nettement entre les années, particulièrement pour Conforme et PIETRA.")
    add_lead(doc, "Hypothèses non tranchées",
             "Évolution des lots, conditions de production, acquisition, diagnostic ou protocole de labellisation.")
    add_lead(doc, "Implication", "L’année devra être contrôlée dans la construction des futurs splits.")
    add_picture(doc, assets["acquisition"],
                "Figure 4 — Équilibre global des emplacements caméra et des positions.", 15.5)
    add_table(doc, ["Caméra", "1", "2", "3", "4", "5", "6"], [[
        "Images", "5 811", "5 953", "6 157", "5 792", "5 913", "5 628"
    ]])
    doc.add_paragraph(
        "Les caméras sont relativement homogènes (15,96 % à 17,46 %). Top représente 17 065 images (48,41 %) "
        "et Bottom 18 189 (51,59 %). Cet équilibre global n’autorise pas à supposer que toutes les images forment "
        "naturellement des paires : il existe 1 124 vues Bottom de plus."
    )

    heading(doc, "2.5 Cas difficiles et couverture de labels_masked", 2)
    add_table(doc, ["Signal", "Images", "Interprétation retenue"], [
        ["multiple=True", "1 197", "Plusieurs fruits visibles ; difficulté réelle"],
        ["chunk=True", "4 975", "Morceau ou débris ; difficulté réelle"],
        ["mixed_quality=True", "0", "Aucune occurrence True dans le fichier fourni"],
        ["labels_masked couvert", "6 577 (18,66 %)", "Source partielle, non vérité globale"],
        ["Conflit vide/non-vide", "59 (0,17 %)", "Sous-ensemble à revoir humainement"],
    ])
    add_lead(doc, "Critique",
             "L’absence de mixed_quality=True peut signifier l’absence du phénomène ou l’absence d’usage du champ. "
             "Le dataset ne permet pas de distinguer ces deux explications.")
    add_table(doc, ["masked_label", "Images", "% des 6 577 couvertes"], [
        ["chataigne", "4 190", "63,71 %"], ["vide", "1 572", "23,90 %"],
        ["chunks", "586", "8,91 %"], ["multiple", "229", "3,48 %"],
    ])

    heading(doc, "3. Extrait vidéo : produire une référence locale", 1)
    heading(doc, "3.1 Décodage et crops manuels", 2)
    add_table(doc, ["Caractéristique", "Top", "Bottom"], [
        ["Durée", "30 s", "30 s"], ["Frames décodées", "750", "750"],
        ["Cadence", "25 FPS", "25 FPS"], ["Résolution", "720 × 576", "720 × 576"],
        ["Codec", "dvsd", "dvsd"], ["Frames retenues", "25", "26"],
        ["Crops", "26", "26"],
    ])
    add_lead(doc, "Choix", "Recadrage manuel : fruit entier avec une petite marge.")
    add_lead(doc, "Intérêt",
             "Le volume est faible, le sujet demande une extraction manuelle et aucun détecteur automatique n’est "
             "introduit avant de construire la référence.")
    add_lead(doc, "Limites",
             "Cadrage subjectif, variabilité humaine, faible scalabilité et sélection potentiellement non exhaustive.")
    doc.add_paragraph(
        "Les nouvelles frames, crops et métadonnées restent locales au dépôt Git du rendu et ne sont pas ajoutées au "
        "Drive DVC partagé. Les 52 crops ne sont pas intégrés au manifeste historique à quatre classes : le mot "
        "« conforme » dans le nom des AVI ne constitue pas un label."
    )

    heading(doc, "3.2 Appariement Top/Bottom par temporalité", 2)
    add_lead(doc, "Problème",
             "Les vues Top et Bottom montrent des faces différentes : une ressemblance visuelle ne peut pas prouver "
             "qu’il s’agit du même fruit.")
    add_lead(doc, "Hypothèse", "Sur cet extrait, la vue Bottom apparaît environ sept frames après la vue Top.")
    add_callout(doc, "Fonction de coût",
                "cost(T,B) = abs((frame_B − frame_T) − 7), avec alignement monotone et gaps autorisés.")
    add_picture(doc, assets["video"],
                "Figure 5 — Résultats de l’alignement temporel sur les 52 crops.", 15.5)
    add_table(doc, ["Métrique", "Résultat"], [
        ["matched_temporal", "22"], ["unmatched_top", "2"], ["unmatched_bottom", "3"],
        ["groupe ambiguous", "1"], ["delta moyen", "6,773 frames"], ["médiane", "7 frames / 280 ms"],
        ["écart-type", "0,419 frame"], ["couverture certaine", "84,62 % Top ; 84,62 % Bottom"],
    ])
    add_lead(doc, "Point satisfaisant",
             "Les gaps évitent un décalage cumulatif et 17 paires sur 22 ont exactement +7 frames.")
    add_lead(doc, "Limite",
             "Trente secondes ne suffisent pas à établir +7 comme règle universelle de la machine.")

    heading(doc, "3.3 L’ambiguïté comme résultat valide", 2)
    add_picture(doc, assets["ambiguous"],
                "Figure 6 — Cas ambigu : deux crops Top au frame 572 pour un crop Bottom au frame 579.", 15.8)
    doc.add_paragraph(
        "T_frame_000572_crop_01 et T_frame_000572_crop_02 donnent le même coût pour "
        "B_frame_000579_crop_01. Aucun candidat n’est choisi arbitrairement. Une méthode de qualité doit pouvoir "
        "exprimer « indéterminé » plutôt que fabriquer une vérité terrain."
    )

    heading(doc, "4. Reconstruction T/B du dataset historique", 1)
    heading(doc, "4.1 Pourquoi la règle vidéo ne se transpose pas", 2)
    doc.add_paragraph(
        "Les 35 254 images historiques n’ont pas de timestamp exploitable. Le décalage vidéo ne peut donc pas être "
        "appliqué. L’audit s’appuie sur les composantes du filename, sans fusionner les labels des deux vues."
    )
    add_table(doc, ["Clé", "Définition"], [
        ["A", "year + batch_num + cam_num + sample_num"],
        ["B", "A + label_filename"],
        ["C", "year + cam_num + sample_num + label_filename"],
        ["D", "year + cam_num + sample_num"],
    ])
    add_picture(doc, assets["keys"],
                "Figure 7 — Couverture obtenue par les quatre clés historiques candidates.", 14.8)
    add_table(doc, ["Clé", "Groupes 1T/1B", "Couverture images"], [
        ["A", "5 561", "31,55 %"], ["B", "16 176", "91,77 %"],
        ["C", "12 318", "69,88 %"], ["D", "932", "5,29 %"],
    ])

    heading(doc, "4.2 Justification et qualification de la clé B", 2)
    add_lead(doc, "Observations",
             "sample_num est massivement réutilisé entre labels historiques : 62,38 % des images appartiennent à "
             "des slots réutilisés, et 86,72 % en 2026. batch_num résout en outre des collisions propres à 2025.")
    add_lead(doc, "Cohérence forte",
             "Les 16 176 groupes B à 1T/1B ont des filenames identiques dans toutes leurs composantes, sauf Cam_T/Cam_B.")
    add_lead(doc, "Limite fondamentale",
             "Aucun générateur historique ni identifiant physique n’a été retrouvé. Une excellente cohérence de nommage "
             "ne prouve pas l’identité physique du fruit.")
    add_callout(doc, "Décision",
                "La clé B est retenue comme fortement soutenue mais restant heuristique : "
                "year + batch_num + cam_num + sample_num + label_filename.")
    doc.add_paragraph(
        "label_filename n’est ni une cible ni une caractéristique ML. Il sert uniquement à interpréter le namespace "
        "historique de sample_num. La cible reste label_principal, puis effective_label après adjudication."
    )

    heading(doc, "4.3 Résultat et sémantique", 2)
    add_table(doc, ["Indicateur", "Avant adjudication", "Après adjudication"], [
        ["Paires heuristiques", "16 176", "16 176"],
        ["Images appariées", "32 352", "32 352"],
        ["Images non appariées", "2 902", "2 902"],
        ["Paires même label", "10 714", "10 721"],
        ["Paires labels différents", "5 462", "5 455"],
    ])
    doc.add_paragraph(
        "Parmi les non appariées figurent 889 Top et 2 013 Bottom. Les labels des deux vues restent indépendants : "
        "une paire technique n’est jamais transformée en une étiquette unique de châtaigne."
    )

    heading(doc, "5. Arbitrer les ambiguïtés sans déformer la donnée", 1)
    heading(doc, "5.1 Comparaison de stratégies", 2)
    add_table(doc, ["Stratégie", "Règle", "Images conservées"], [
        ["S0", "Aucun filtre", "35 254"],
        ["S1", "Isoler les 59 conflits masked/principal", "35 195"],
        ["S2 / S3a", "Ajouter multiple, chunk, mixed_quality", "29 043"],
        ["S3b", "Ajouter l’exigence de provenance", "22 103"],
    ])
    add_picture(doc, assets["strategies"],
                "Figure 8 — Effet des stratégies sur le volume et la composition des classes.", 15.8)
    add_lead(doc, "Constat",
             "S2 fait passer Vide de 36,10 % à 43,66 %, Conforme de 30,74 % à 27,28 % et PIETRA de 17,11 % à 13,00 %. "
             "S3b accentue encore les biais.")
    add_lead(doc, "Décision",
             "multiple, chunk et les 8 033 lignes sans labeled_by sont conservés. Absence de provenance ne signifie pas "
             "mauvais label ; les difficultés observées font partie de la production.")
    add_lead(doc, "Critique",
             "Ces cas compliqueront l’apprentissage. Les flags sont donc maintenus pour analyser les erreurs futures.")

    heading(doc, "5.2 Revue humaine ciblée des 59 conflits", 2)
    doc.add_paragraph(
        "Les 59 conflits vide/non-vide représentent 0,17 % du dataset. Une revue en aveugle était moins coûteuse "
        "qu’une suppression automatique et évitait de privilégier arbitrairement label_principal ou labels_masked. "
        "Avant chaque décision, aucune annotation ni métrique existante n’était affichée."
    )
    add_table(doc, ["Résultat de la revue binaire", "Nombre"], [
        ["Vide", "45"], ["NonVide", "14"], ["Incertain", "0"],
        ["Accord humain / principal", "43"], ["Accord humain / masked", "16"],
    ])
    add_callout(doc, "Précaution d’interprétation",
                "Le ratio 43/59 ne mesure pas la fiabilité globale de label_principal : les 59 images sont un "
                "sous-ensemble pré-sélectionné précisément parce que les sources étaient en conflit.", "FFF4E2")

    heading(doc, "5.3 Seconde revue multiclasses", 2)
    doc.add_paragraph(
        "Neuf images jugées NonVide avaient label_principal=Vide. La décision binaire ne pouvait pas déterminer "
        "Conforme, NON Conforme ou PIETRA. Une seconde revue multiclasses, toujours en aveugle, a produit "
        "3 Conforme, 5 NON Conforme et 1 PIETRA, sans Incertain. Cette organisation limite l’effort complexe à neuf cas."
    )
    add_table(doc, ["Repère opérationnel de revue", "Description pratique"], [
        ["Conforme", "Fruit clair et sain visuellement, sans moisissure ni défaut marqué"],
        ["NON Conforme", "Fruit très noir, fortement dégradé, éclaté ou avec défauts importants"],
        ["PIETRA", "État intermédiaire"],
    ])
    doc.add_paragraph(
        "Ces repères ont guidé la revue humaine ; ce ne sont pas des définitions réglementaires officielles extraites de l’annexe.",
        style="Small note",
    )

    heading(doc, "6. Dataset final et traçabilité", 1)
    heading(doc, "6.1 Décision finale", 2)
    add_picture(doc, assets["final"],
                "Figure 9 — Distribution finale et provenance de effective_label.", 15.8)
    add_table(doc, ["effective_label", "Images", "Part", "Écart vs label_principal"], [
        ["Vide", "12 724", "36,09 %", "−2"],
        ["Conforme", "10 832", "30,73 %", "−4"],
        ["PIETRA", "6 034", "17,12 %", "+1"],
        ["NON Conforme", "5 664", "16,07 %", "+5"],
    ])
    add_lead(doc, "Résultat", "Les 35 254 images sont conservées ; 16 labels seulement sont effectivement corrigés.")
    add_table(doc, ["Transition effective", "Images"], [
        ["Conforme → Vide", "7"], ["Vide → Conforme", "3"],
        ["Vide → NON Conforme", "5"], ["Vide → PIETRA", "1"],
        ["Sans changement", "35 238"],
    ])
    add_lead(doc, "Implication",
             "La stabilité globale confirme une correction ciblée, non une reconstruction massive des labels.")

    heading(doc, "6.2 Provenance du label final", 2)
    add_table(doc, ["effective_label_source", "Images", "Part"], [
        ["principal", "35 195", "99,83 %"],
        ["principal_confirmed_by_human", "43", "0,12 %"],
        ["human_binary_review", "7", "0,02 %"],
        ["human_multiclass_review", "9", "0,03 %"],
    ])
    doc.add_paragraph(
        "La correction n’écrase aucune source. label_principal et les signaux masked restent présents, tandis que "
        "effective_label_source et adjudication_status expliquent chaque décision."
    )

    heading(doc, "6.3 Manifeste ML et protection contre la fuite", 2)
    add_callout(doc, "Référence du §4.2", "analysis/output/final_training_manifest.csv")
    doc.add_paragraph(
        "Le manifeste dérivé contient 35 254 filenames uniques, effective_label, les métadonnées d’acquisition, les "
        "flags de difficulté, la provenance des labels, historical_pair_id et split_group_id. Pour les paires "
        "historiques, split_group_id est commun aux deux vues ; chaque image non appariée reçoit un groupe déterministe propre."
    )
    add_lead(doc, "Règle future",
             "Au §4.2, les groupes devront rester indivisibles afin d’éviter qu’un Top soit placé en train et son Bottom potentiel en test.")

    heading(doc, "6.4 Validation technique", 2)
    add_table(doc, ["Famille de contrôle", "Résultat"], [
        ["Volume et unicité", "35 254 lignes ; 35 254 filenames uniques"],
        ["Cible", "35 254 labels ; 4 classes exactes ; aucune valeur vide"],
        ["Adjudication", "0 pending ; 0 uncertain"],
        ["Paires", "32 352 images ; 16 176 pair_id ; deux membres par paire"],
        ["Non appariées", "2 902 ; un split_group_id propre par image"],
        ["Cohérence groupes", "Membres d’une paire groupés ; groupes tous disjoints"],
        ["Bilan", "19 / 19 contrôles réussis"],
    ])

    heading(doc, "7. Proposition de processus collaboratif", 1)
    doc.add_paragraph(
        "L’objectif n’est pas de bâtir une plateforme lourde, mais de supprimer les ambiguïtés observées sans "
        "recréer les conflits d’un gros CSV modifié par plusieurs personnes."
    )
    add_table(doc, ["Problème observé", "Évolution minimale", "Bénéfice"], [
        ["Pseudos libres", "Utilisateur authentifié ou liste contrôlée", "Identités stables"],
        ["reviewed sans validateur", "Interdire reviewed=True sans validateur", "Statut vérifiable"],
        ["Écrasement de label", "Journal ancien/nouveau/auteur/date/raison", "Audit complet"],
        ["Conflits Git sur le CSV", "Streamlit + SQLite locale/partagée ou petite base centrale", "Écritures atomiques"],
        ["Déploiement sans base", "Un fichier par annotateur/lot puis fusion contrôlée", "Alternative simple"],
        ["Cas ambigus", "uncertain, multiple, chunk, mixed_quality + file de seconde revue", "Incertitude explicite"],
        ["Contrôle qualité", "Double validation ciblée + échantillon aléatoire", "Effort proportionné"],
    ], font_size=7.8)
    add_lead(doc, "Décision",
             "Ne pas doubler 35 000 annotations : concentrer la seconde lecture sur les conflits, les uncertain et un échantillon aléatoire.")

    heading(doc, "8. Limites et regard critique", 1)
    add_table(doc, ["Satisfaisant", "Moins satisfaisant / risque ouvert"], [
        ["35 254 filenames uniques", "Absence d’identifiant physique de châtaigne"],
        ["100 % des images conservées", "Provenance incomplète pour 8 033 lignes"],
        ["16 corrections finales ciblées", "Contradiction entre reviewed et le sujet"],
        ["91,77 % de couverture T/B historique", "Appariement historique heuristique"],
        ["Temporalité vidéo stable sur l’extrait", "30 secondes ; +7 non généralisable sans validation"],
        ["Corrections et sources traçables", "Grille visuelle subjective et non réglementaire"],
        ["Manifeste reproductible", "mixed_quality sans occurrence True"],
        ["Difficultés réelles conservées", "Évolution marquée entre 2025 et 2026"],
    ], widths=[8, 8], font_size=8.1)
    doc.add_paragraph(
        "Les annotations restent historiques et humaines. Certaines vues appariées portent des labels différents, "
        "sans qu’il soit légitime de les fusionner. Le générateur de filenames n’a pas été retrouvé. Les crops vidéo "
        "sont manuels et potentiellement non exhaustifs. Enfin, conserver chunk et multiple maintient une difficulté "
        "réaliste, mais demandera une analyse d’erreurs dédiée lors de l’évaluation."
    )
    add_callout(doc, "Position critique",
                "Le dataset n’est pas « propre » au sens d’une vérité parfaite ; il est désormais explicite, traçable "
                "et structuré pour que ses incertitudes ne soient pas confondues avec des certitudes.")

    heading(doc, "9. Conclusion du §4.1", 1)
    add_table(doc, ["Question", "Conclusion"], [
        ["Dataset exploitable ?", "Oui, avec les précautions documentées"],
        ["Dataset final ?", "35 254 images avec effective_label"],
        ["Corrections ?", "16 transitions humaines effectives"],
        ["Données difficiles ?", "Conservées car représentatives de la production"],
        ["Top/Bottom vidéo ?", "Alignement temporel, gaps et ambiguïtés explicites"],
        ["Top/Bottom historique ?", "Namespace fortement soutenu mais heuristique"],
        ["Limite principale ?", "Pas de vérité physique T/B ; hétérogénéité historique"],
        ["Préparation du §4.2 ?", "split_group_id prêt pour des splits sans fuite entre vues"],
    ])
    doc.add_paragraph(
        "Le §4.1 aboutit donc à une référence ML dérivée sans altération silencieuse des sources. L’étape suivante "
        "pourra construire des splits contrôlés, mais elle devra préserver les groupes T/B, surveiller la dérive annuelle "
        "et évaluer prioritairement la pureté et la complétude de la classe Conforme conformément à l’annexe métier."
    )

    heading(doc, "Annexe A — Comparatif technique des clés T/B", 1)
    add_table(doc, ["Clé", "Composantes", "1T/1B", "Couverture", "Lecture"], [
        ["A", "year, batch_num, cam_num, sample_num", "5 561", "31,55 %", "Collisions entre labels"],
        ["B", "A + label_filename", "16 176", "91,77 %", "Retenue, heuristique"],
        ["C", "year, cam_num, sample_num, label_filename", "12 318", "69,88 %", "Perd batch_num 2025"],
        ["D", "year, cam_num, sample_num", "932", "5,29 %", "Collisions massives"],
    ], font_size=7.7)
    doc.add_paragraph(
        "62,38 % des images appartiennent à des slots sample_num réutilisés entre labels historiques ; le taux atteint "
        "86,72 % en 2026. Le fait que 16 176/16 176 paires B ne diffèrent que par Cam_T/Cam_B soutient fortement "
        "la clé, sans fournir d’identifiant physique."
    )

    heading(doc, "Annexe B — Tables de distribution", 1)
    add_table(doc, ["Dimension", "Catégorie", "Nombre", "Pourcentage"], [
        ["Année", "2025", "20 266", "57,49 %"], ["Année", "2026", "14 988", "42,51 %"],
        ["Position", "Top", "17 065", "48,41 %"], ["Position", "Bottom", "18 189", "51,59 %"],
        ["Caméra", "1", "5 811", "16,48 %"], ["Caméra", "2", "5 953", "16,89 %"],
        ["Caméra", "3", "6 157", "17,46 %"], ["Caméra", "4", "5 792", "16,43 %"],
        ["Caméra", "5", "5 913", "16,77 %"], ["Caméra", "6", "5 628", "15,96 %"],
    ])
    add_table(doc, ["Label", "Initial", "Final", "Écart"], [
        ["Vide", "12 726", "12 724", "−2"], ["Conforme", "10 836", "10 832", "−4"],
        ["PIETRA", "6 033", "6 034", "+1"], ["NON Conforme", "5 659", "5 664", "+5"],
    ])

    heading(doc, "Annexe C — Traçabilité technique", 1)
    add_table(doc, ["Artefact", "Rôle"], [
        ["analysis/build_dataset_manifest.py", "Manifeste d’audit et jointure masked non destructive"],
        ["analysis/dataset_quality.py", "Statistiques qualité et exports d’inspection"],
        ["analysis/video_manual_crop.py", "Extraction et recadrage manuel"],
        ["analysis/tb_pairing_tool.py", "Audit et alignement temporel de la vidéo"],
        ["analysis/historical_tb_audit.py", "Comparaison des namespaces historiques"],
        ["analysis/historical_tb_pairing.py", "Table des paires heuristiques historiques"],
        ["analysis/review_masked_disagreements.py", "Première revue en aveugle"],
        ["analysis/review_multiclass_conflicts.py", "Seconde revue multiclasses en aveugle"],
        ["analysis/adjudicate_masked_disagreements.py", "Adjudication dérivée et traçable"],
        ["analysis/final_dataset_strategy.py", "Simulation des stratégies d’inclusion"],
        ["analysis/build_final_training_manifest.py", "Manifeste final et validations"],
    ], font_size=7.5)
    doc.add_paragraph(
        "Sources documentaires : [1] 25-26_eval_ipg_bihar_castagnet_v1.pdf, §4.1 ; "
        "[2] 25-26_eval_ipg_bihar_castagnet-annexea_v1.pdf, exigences qualité GRPTMC ; "
        "[3] Rapport_Dataset_Chataignes.pdf, inventaire historique fondé sur le filename.", style="Small note")

    heading(doc, "Annexe D — Checklist de couverture du §4.1", 1)
    requirements = [
        "Extraction vidéo", "Crop manuel", "Correspondance T/B vidéo", "Correspondance T/B historique",
        "Analyse qualité", "Distribution des labels", "Distribution par année", "Distribution par caméra",
        "Taux de relecture", "Cas ambigus", "Hétérogénéité du diagnostic", "Décisions inclusion/exclusion",
        "Revue humaine", "Dataset final", "Proposition collaborative", "Critique", "Limites",
    ]
    add_table(doc, ["Exigence", "Couverture", "Section"], [
        [req, "Couverte", section] for req, section in zip(requirements, [
            "3.1", "3.1", "3.2–3.3", "4", "2", "2.1 / 6.1", "2.4", "2.4",
            "2.3", "2.5 / 3.3 / 5", "2.1 / 2.4", "5.1", "5.2–5.3", "6",
            "7", "8", "8",
        ])
    ], font_size=7.5)
    doc.add_paragraph(
        "Périmètre volontaire : ce document s’arrête à la constitution et à la validation de la donnée. Il ne crée "
        "aucun split, n’entraîne aucun modèle et ne présente aucun résultat du §4.2.", style="Small note")
    return doc


def main() -> None:
    REPORTS.mkdir(parents=True, exist_ok=True)
    doc = build_report()
    doc.save(OUTPUT)
    print(f"Rapport créé : {OUTPUT}")
    print("Figures : 9")
    print("Contrôles du manifeste final : 19/19 PASS")


if __name__ == "__main__":
    main()
