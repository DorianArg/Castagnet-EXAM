"""Génère le rapport §4.2 exclusivement depuis les artefacts existants."""

from __future__ import annotations

import json
import sqlite3
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "analysis" / "output"
MODELING = OUTPUT / "modeling"
FIGURES = OUTPUT / "report_4_2_figures"
REPORTS = ROOT / "reports"
DOCX_PATH = REPORTS / "Rapport_CastagNet_4_2_Modelisation.docx"
MD_PATH = REPORTS / "Rapport_CastagNet_4_2_Modelisation.md"
MLFLOW_DB = ROOT / "mlflow.db"

NAVY = "17324D"
BLUE = "245B78"
TEAL = "2A7F7F"
PALE_BLUE = "EAF2F7"
PALE_TEAL = "E8F3F1"
PALE_RED = "FBEAEA"
PALE_AMBER = "FFF4DC"
DARK = "263238"
MID = "5D6A72"
WHITE = "FFFFFF"


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def pct(value: float, digits: int = 2) -> str:
    return f"{value * 100:.{digits}f} %".replace(".", ",")


def integer(value: int) -> str:
    return f"{value:,}".replace(",", " ")


def duration(seconds: float) -> str:
    total = int(round(seconds))
    hours, rest = divmod(total, 3600)
    minutes, secs = divmod(rest, 60)
    return f"{hours} h {minutes:02d} min {secs:02d} s"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, end))


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, color, before, after in (
        ("Title", 29, NAVY, 0, 12),
        ("Subtitle", 15, TEAL, 0, 10),
        ("Heading 1", 18, NAVY, 16, 8),
        ("Heading 2", 14, BLUE, 12, 6),
        ("Heading 3", 11.5, TEAL, 9, 4),
    ):
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    if "Caption Report" not in [style.name for style in document.styles]:
        caption = document.styles.add_style("Caption Report", WD_STYLE_TYPE.PARAGRAPH)
        caption.font.name = "Aptos"
        caption.font.size = Pt(8.5)
        caption.font.italic = True
        caption.font.color.rgb = RGBColor.from_string(MID)
        caption.paragraph_format.space_before = Pt(3)
        caption.paragraph_format.space_after = Pt(8)
    header = section.header.paragraphs[0]
    header.text = "CASTAGNET   |   §4.2 Modélisation et évaluation"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Aptos"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string(MID)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Rapport final — 3 septembre 2026    •    ")
    add_field(footer, "PAGE")
    for run in footer.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MID)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


class Report:
    def __init__(self) -> None:
        self.doc = Document()
        configure_document(self.doc)
        self.md: list[str] = []
        self.table_number = 0
        self.figure_number = 0
        self.figure_titles: list[str] = []
        self.table_titles: list[str] = []

    def title_page(self) -> None:
        self.md.extend([
            "# CASTAGNET — §4.2 Modélisation et évaluation",
            "",
            "**Rapport final — Modélisation, entraînement, comparaison et évaluation**",
            "",
            "Mise en Situation Professionnelle — MSc BIHAR  ",
            "Compétences C28 et C31  ",
            "3 septembre 2026",
            "",
        ])
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(52)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("CASTAGNET")
        run.bold = True
        run.font.name = "Aptos Display"
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor.from_string(TEAL)
        p = self.doc.add_paragraph(style="Title")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run("Modélisation et évaluation\ndes modèles de classification")
        p = self.doc.add_paragraph(style="Subtitle")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run("Rapport final — §4.2")
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(30)
        run = p.add_run("Mise en Situation Professionnelle — MSc BIHAR\nCompétences C28 et C31\n3 septembre 2026")
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor.from_string(MID)
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(55)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("Configuration finale figée avant test\nResNet18 • checkpoint epoch 13 • seuil Conforme 0,55")
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor.from_string(NAVY)
        self.doc.add_page_break()

    def heading(self, text: str, level: int = 1, page_break: bool = False) -> None:
        if page_break:
            self.doc.add_page_break()
        self.doc.add_heading(text, level=level)
        self.md.extend(["#" * (level + 1) + " " + text, ""])

    def paragraph(self, text: str) -> None:
        self.doc.add_paragraph(text)
        self.md.extend([text, ""])

    def bullets(self, items: list[str]) -> None:
        for item in items:
            p = self.doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2.5)
            p.add_run(item)
            self.md.append(f"- {item}")
        self.md.append("")

    def callout(self, title: str, text: str, kind: str = "info") -> None:
        fill = {"info": PALE_BLUE, "success": PALE_TEAL, "warning": PALE_AMBER, "risk": PALE_RED}[kind]
        table = self.doc.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        set_cell_shading(cell, fill)
        set_cell_margins(cell, 140, 160, 140, 160)
        p = cell.paragraphs[0]
        run = p.add_run(title + " — ")
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(NAVY if kind != "risk" else "8B1E1E")
        p.add_run(text)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(1)
        self.md.extend([f"> **{title} —** {text}", ""])

    def table(self, title: str, headers: list[str], rows: list[list[str]], source: str, widths=None) -> None:
        self.table_number += 1
        full_title = f"Tableau {self.table_number} — {title}"
        self.table_titles.append(full_title)
        p = self.doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        run = p.add_run(full_title)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor.from_string(NAVY)
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.autofit = True
        header = table.rows[0]
        set_repeat_table_header(header)
        for index, value in enumerate(headers):
            cell = header.cells[index]
            cell.text = value
            set_cell_shading(cell, NAVY)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for cell_run in cell.paragraphs[0].runs:
                cell_run.bold = True
                cell_run.font.color.rgb = RGBColor.from_string(WHITE)
                cell_run.font.size = Pt(8.2)
        for row_index, values in enumerate(rows):
            cells = table.add_row().cells
            prevent_row_split(table.rows[-1])
            if row_index % 2:
                for cell in cells:
                    set_cell_shading(cell, "F5F8FA")
            for index, value in enumerate(values):
                cells[index].text = str(value)
                cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cells[index])
                for cell_run in cells[index].paragraphs[0].runs:
                    cell_run.font.size = Pt(7.8)
        if widths:
            for row in table.rows:
                for index, width in enumerate(widths):
                    row.cells[index].width = Cm(width)
        caption = self.doc.add_paragraph(style="Caption Report")
        caption.add_run(f"Source : {source}")
        self.md.extend([f"**{full_title}**", "", "| " + " | ".join(headers) + " |", "|" + "|".join(["---"] * len(headers)) + "|"])
        for values in rows:
            self.md.append("| " + " | ".join(str(value) for value in values) + " |")
        self.md.extend(["", f"*Source : {source}*", ""])

    def figure(self, title: str, path: Path, source: str, width_cm: float = 15.8) -> None:
        self.figure_number += 1
        full_title = f"Figure {self.figure_number} — {title}"
        self.figure_titles.append(full_title)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        caption = self.doc.add_paragraph(style="Caption Report")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.add_run(f"{full_title}\nSource : {source}")
        self.md.extend([f"![{full_title}]({path.relative_to(ROOT).as_posix()})", "", f"*{full_title}. Source : {source}*", ""])

    def save(self) -> None:
        self.doc.core_properties.title = "CastagNet — §4.2 Modélisation et évaluation"
        self.doc.core_properties.subject = "Rapport final de modélisation CastagNet"
        self.doc.core_properties.author = "Projet CastagNet — MSc BIHAR"
        self.doc.core_properties.keywords = "CastagNet, classification, ResNet18, MLflow, validation, test"
        update_fields = OxmlElement("w:updateFields")
        update_fields.set(qn("w:val"), "true")
        self.doc.settings._element.append(update_fields)
        REPORTS.mkdir(parents=True, exist_ok=True)
        self.doc.save(DOCX_PATH)
        MD_PATH.write_text("\n".join(self.md), encoding="utf-8")


def generate_figures(histories: dict[str, pd.DataFrame], calibration: pd.DataFrame, test: dict) -> dict[str, Path]:
    FIGURES.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({
        "font.family": "DejaVu Sans",
        "font.size": 9,
        "axes.titlesize": 12,
        "axes.labelsize": 9,
        "axes.spines.top": False,
        "axes.spines.right": False,
        "axes.grid": True,
        "grid.alpha": .18,
        "figure.facecolor": "white",
    })
    outputs: dict[str, Path] = {}
    fig, axes = plt.subplots(1, 2, figsize=(11, 4.2))
    for ax, key, title in zip(axes, ("custom_v1", "custom_v2"), ("CustomCNN v1", "CustomCNN v2")):
        frame = histories[key]
        ax.plot(frame.epoch, frame.train_loss, marker="o", color="#2A7F7F", label="train_loss")
        ax.plot(frame.epoch, frame.val_loss, marker="o", color="#C94C4C", label="val_loss")
        ax.set_title(title)
        ax.set_xlabel("Epoch")
        ax.set_ylabel("Loss")
        ax.legend(frameon=False)
    fig.suptitle("CNN maison — dynamique des losses", fontweight="bold", color="#17324D")
    fig.tight_layout()
    outputs["custom_loss"] = FIGURES / "customcnn_losses.png"
    fig.savefig(outputs["custom_loss"], dpi=190, bbox_inches="tight")
    plt.close(fig)
    resnet = histories["resnet18"]
    for key, train_col, val_col, ylabel, filename in (
        ("resnet_loss", "train_loss", "val_loss", "Loss", "resnet18_losses.png"),
        ("resnet_accuracy", "train_accuracy", "val_accuracy", "Accuracy", "resnet18_accuracy.png"),
    ):
        fig, ax = plt.subplots(figsize=(9.3, 4.5))
        ax.plot(resnet.epoch, resnet[train_col], marker="o", color="#2A7F7F", label=train_col)
        ax.plot(resnet.epoch, resnet[val_col], marker="o", color="#C94C4C", label=val_col)
        ax.axvspan(.5, 3.5, color="#EAF2F7", alpha=.9, label="tête seule")
        ax.axvline(13, color="#17324D", linestyle="--", linewidth=1.3, label="meilleur checkpoint")
        ax.set_xlabel("Epoch globale")
        ax.set_ylabel(ylabel)
        ax.set_xticks(resnet.epoch)
        ax.legend(frameon=False, ncol=2)
        ax.set_title(f"ResNet18 — {ylabel.lower()} train / validation", fontweight="bold", color="#17324D")
        fig.tight_layout()
        outputs[key] = FIGURES / filename
        fig.savefig(outputs[key], dpi=190, bbox_inches="tight")
        plt.close(fig)
    fig, ax = plt.subplots(figsize=(9.3, 4.5))
    ax.plot(resnet.epoch, resnet.val_macro_f1, marker="o", color="#245B78", label="macro-F1 validation")
    ax.axvspan(.5, 3.5, color="#EAF2F7", alpha=.9, label="tête seule")
    ax.axvline(13, color="#17324D", linestyle="--", linewidth=1.3, label="meilleur checkpoint")
    ax.scatter([13], [resnet.loc[resnet.epoch == 13, "val_macro_f1"].iloc[0]], color="#C94C4C", zorder=4)
    ax.set_xlabel("Epoch globale")
    ax.set_ylabel("Macro-F1")
    ax.set_xticks(resnet.epoch)
    ax.legend(frameon=False)
    ax.set_title("ResNet18 — macro-F1 de validation", fontweight="bold", color="#17324D")
    fig.tight_layout()
    outputs["resnet_f1"] = FIGURES / "resnet18_val_macro_f1.png"
    fig.savefig(outputs["resnet_f1"], dpi=190, bbox_inches="tight")
    plt.close(fig)
    fig, ax = plt.subplots(figsize=(9.3, 4.8))
    ax.plot(calibration.threshold, calibration.precision_conforme, color="#245B78", label="Précision Conforme")
    ax.plot(calibration.threshold, calibration.recall_conforme, color="#2A7F7F", label="Rappel Conforme")
    ax.axhline(.95, color="#245B78", linestyle=":", label="objectif précision 95 %")
    ax.axhline(.85, color="#2A7F7F", linestyle=":", label="objectif rappel 85 %")
    for threshold, color, label in ((.55, "#17324D", "seuil retenu 0,55"), (.56, "#C98A00", "compromis 0,56"), (.91, "#A33A3A", "premier seuil ≥95 %")):
        ax.axvline(threshold, color=color, linestyle="--", linewidth=1.1, label=label)
    ax.set_xlabel("Seuil P(Conforme)")
    ax.set_ylabel("Score")
    ax.set_ylim(.55, 1.01)
    ax.legend(frameon=False, ncol=2, fontsize=8)
    ax.set_title("Calibration du seuil Conforme — validation uniquement", fontweight="bold", color="#17324D")
    fig.tight_layout()
    outputs["calibration"] = FIGURES / "resnet18_conforme_threshold_calibration.png"
    fig.savefig(outputs["calibration"], dpi=190, bbox_inches="tight")
    plt.close(fig)
    labels = ["Conforme", "NON Conforme", "PIETRA", "Vide"]
    for key, matrix_key, title, filename in (
        ("test_argmax", "argmax", "Test final — décision argmax", "test_confusion_argmax.png"),
        ("test_threshold", "threshold_055", "Test final — seuil Conforme 0,55", "test_confusion_threshold_055.png"),
    ):
        matrix = np.array(test[matrix_key]["confusion_matrix"])
        fig, ax = plt.subplots(figsize=(6.4, 5.4))
        image = ax.imshow(matrix, cmap="Blues")
        cutoff = matrix.max() * .52
        for row in range(4):
            for column in range(4):
                ax.text(column, row, str(matrix[row, column]), ha="center", va="center", color="white" if matrix[row, column] > cutoff else "#" + DARK)
        ax.set_xticks(range(4), labels, rotation=25, ha="right")
        ax.set_yticks(range(4), labels)
        ax.set_xlabel("Classe prédite")
        ax.set_ylabel("Classe réelle")
        ax.set_title(title, fontweight="bold", color="#17324D")
        fig.colorbar(image, ax=ax, fraction=.046)
        fig.tight_layout()
        outputs[key] = FIGURES / filename
        fig.savefig(outputs[key], dpi=190, bbox_inches="tight")
        plt.close(fig)
    return outputs


def verify_mlflow(run_ids: dict[str, str]) -> dict[str, dict[str, object]]:
    connection = sqlite3.connect(MLFLOW_DB)
    tables = {row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    metric_table = "latest_metrics" if "latest_metrics" in tables else "metrics"
    verified: dict[str, dict[str, object]] = {}
    for model, run_id in run_ids.items():
        row = connection.execute(
            "SELECT status, start_time, end_time FROM runs WHERE run_uuid = ?", (run_id,)
        ).fetchone()
        if row is None:
            raise RuntimeError(f"Run MLflow absent : {model} / {run_id}")
        metrics = dict(connection.execute(
            f"SELECT key, value FROM {metric_table} WHERE run_uuid = ?", (run_id,)
        ).fetchall())
        verified[model] = {
            "run_id": run_id,
            "status": row[0],
            "start_time": row[1],
            "end_time": row[2],
            "metrics": metrics,
        }
    connection.close()
    return verified


def build_report() -> tuple[Report, dict[str, object]]:
    split = load_json(OUTPUT / "ml_split_summary.json")
    final_manifest = load_json(OUTPUT / "final_training_manifest_summary.json")
    model_keys = {
        "CustomCNN v1": "custom_cnn_baseline_v1",
        "CustomCNN v2": "custom_cnn_baseline_v2",
        "MobileNetV3 Small": "mobilenet_v3_small",
        "EfficientNet-B0": "efficientnet_b0",
        "ResNet18": "resnet18",
    }
    metrics = {
        name: load_json(MODELING / folder / "metrics_summary.json")
        for name, folder in model_keys.items()
    }
    smoke = {
        name: load_json(MODELING / folder / "smoke_test_summary.json")
        for name, folder in model_keys.items()
    }
    histories = {
        "custom_v1": pd.read_csv(MODELING / "custom_cnn_baseline_v1" / "training_history.csv"),
        "custom_v2": pd.read_csv(MODELING / "custom_cnn_baseline_v2" / "training_history.csv"),
        "resnet18": pd.read_csv(MODELING / "resnet18" / "training_history.csv"),
    }
    calibration = pd.read_csv(MODELING / "resnet18" / "calibration" / "conforme_threshold_curve.csv")
    calibration_summary = load_json(
        MODELING / "resnet18" / "calibration" / "conforme_threshold_summary.json"
    )
    test = load_json(MODELING / "resnet18" / "final_test" / "final_test_metrics.json")
    run_ids = {name: values["run_id"] for name, values in metrics.items()}
    mlflow = verify_mlflow(run_ids)
    figures = generate_figures(histories, calibration, test)
    report = Report()
    report.title_page()

    report.heading("Synthèse décisionnelle", level=1)
    report.callout(
        "Décision",
        "ResNet18 est retenu comme candidat qualité du §4.2. Le checkpoint de l'epoch 13 "
        "et le seuil Conforme 0,55 ont été figés avant l'ouverture du test.",
        "success",
    )
    report.paragraph(
        "Le travail compare une baseline CNN construite depuis zéro à trois architectures "
        "pré-entraînées sur ImageNet. L'emploi du transfer learning améliore nettement la "
        "séparation des classes proches, en particulier PIETRA. Sur validation, ResNet18 "
        f"atteint {pct(metrics['ResNet18']['best_validation']['accuracy'])} d'accuracy et "
        f"{pct(metrics['ResNet18']['best_validation']['macro_f1'])} de macro-F1, contre "
        f"{pct(metrics['CustomCNN v2']['best_validation']['accuracy'])} et "
        f"{pct(metrics['CustomCNN v2']['best_validation']['macro_f1'])} pour la baseline retenue."
    )
    report.paragraph(
        "Le seuil 0,55 rend l'acceptation Conforme plus stricte. Sur le test final unique, "
        f"il réduit le compteur métier de faux Conforme de {test['argmax']['business_errors']['false_conforme_total']} "
        f"à {test['threshold_055']['business_errors']['false_conforme_total']}, mais augmente "
        f"les Conforme rejetées de {test['argmax']['business_errors']['conforme_rejected']} "
        f"à {test['threshold_055']['business_errors']['conforme_rejected']}."
    )
    report.callout(
        "Cahier des charges partiellement atteint",
        f"la configuration finale atteint le rappel minimal ({pct(test['threshold_055']['per_class']['Conforme']['recall'])} "
        f"pour un objectif de 85 %), mais pas la précision ({pct(test['threshold_055']['per_class']['Conforme']['precision'])} "
        "pour un objectif de 95 %). Aucune adaptation post-test n'est autorisée.",
        "warning",
    )
    report.paragraph(
        "Cette sélection reste prédictive, non embarquée. ResNet18 est le plus lourd des trois "
        "modèles pré-entraînés évalués ; sa latence, sa mémoire, son export ONNX et sa capacité "
        "à traiter douze flux sur la cible GTX 1060 3 Go seront mesurés au §4.3."
    )

    report.heading("Sommaire", level=1, page_break=True)
    toc = report.doc.add_paragraph()
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u')
    report.md.extend([
        "1. 4.2.1 Objectifs et contraintes métier",
        "2. 4.2.2 Construction des jeux train, validation et test",
        "3. 4.2.3 Prétraitement et augmentation",
        "4. 4.2.4 Baseline CNN maison",
        "5. 4.2.5 Analyse des deux versions du CNN",
        "6. 4.2.6 Mise en place du transfer learning",
        "7. 4.2.7 Modèles pré-entraînés étudiés",
        "8. 4.2.8 Suivi des expériences avec MLflow",
        "9. 4.2.9 Comparaison des architectures",
        "10. 4.2.10 Sélection de ResNet18",
        "11. 4.2.11 Calibration du seuil Conforme",
        "12. 4.2.12 Évaluation finale sur le jeu de test",
        "13. 4.2.13 Analyse du respect du cahier des charges",
        "14. 4.2.14 Prise en compte des contraintes d'embarquement",
        "15. 4.2.15 Limites, recul critique et pistes d'amélioration",
        "16. 4.2.16 Conclusion et transition vers le §4.3",
        "",
    ])

    report.heading("4.2.1 Objectifs et contraintes métier", level=1, page_break=True)
    report.paragraph(
        "La tâche consiste à classer chaque image dans l'une des quatre classes exclusives : "
        "Conforme, NON Conforme, PIETRA ou Vide. La cible d'apprentissage est exclusivement "
        "effective_label, avec le mapping Conforme=0, NON Conforme=1, PIETRA=2 et Vide=3. "
        "Le nom du fichier et le label historique ne déterminent jamais la classe."
    )
    report.paragraph(
        "Deux contraintes structurent l'évaluation : une précision de la classe Conforme "
        "d'au moins 95 % et un rappel Conforme d'au moins 85 %. La qualité du produit final "
        "est prioritaire sur le rendement : lorsqu'un arbitrage est nécessaire, il est "
        "préférable de rejeter davantage de fruits réellement conformes plutôt que de laisser "
        "passer des fruits NON Conforme ou PIETRA comme Conforme."
    )
    report.callout(
        "Lecture des métriques",
        "la précision Conforme mesure la fiabilité d'une acceptation ; le rappel Conforme "
        "mesure la part des vrais Conforme effectivement acceptés. Les deux objectifs doivent "
        "être satisfaits simultanément.",
        "info",
    )
    report.paragraph(
        "Le contexte de déploiement futur ajoute une contrainte de calcul : NVIDIA GTX 1060 "
        "3 Go, ancien processeur Intel i7, douze flux caméra et export ONNX. Le nombre de "
        "paramètres est donc suivi dès ce chapitre, mais il ne permet pas à lui seul de conclure "
        "sur le débit ou la latence. Cette validation est volontairement reportée au §4.3."
    )

    report.heading("4.2.2 Construction des jeux train, validation et test", level=1)
    report.paragraph(
        f"Le dataset final issu du §4.1 contient {integer(final_manifest['total_images'])} images "
        "uniques. Il a été conservé en entier : réduire arbitrairement ce volume aurait supprimé "
        "de l'information utile sans démontrer que les confusions entre Conforme, PIETRA et "
        "NON Conforme provenaient des données plutôt que du modèle."
    )
    split_rows = []
    for name in ("train", "validation", "test"):
        values = split["splits"][name]
        split_rows.append([
            name.capitalize(),
            integer(values["images"]),
            f"{values['percentage_of_dataset']:.4f} %".replace(".", ","),
            integer(values["split_group_ids"]),
            integer(values["historical_tb"]["complete_pairs"]),
            f"{values['effective_label']['Conforme']['percentage']:.2f} %".replace(".", ","),
            f"{values['effective_label']['PIETRA']['percentage']:.2f} %".replace(".", ","),
        ])
    report.table(
        "Répartition group-aware des données",
        ["Jeu", "Images", "Part", "Groupes", "Paires T/B", "Conforme", "PIETRA"],
        split_rows,
        "analysis/output/ml_split_summary.json",
    )
    report.paragraph(
        "Le ratio 70/15/15 conserve une base d'apprentissage large tout en donnant à la "
        "validation et au test plus de cinq mille images chacun, quantité suffisante pour "
        "produire des métriques stables. Les distributions de classes, années, caméras et "
        "positions ont été maintenues au plus près ; l'écart maximal de proportion de classe "
        f"observé dans le rapport de split n'est que de {split['representativeness']['effective_label']['max_abs_pp_difference']:.4f} point."
    )
    report.paragraph(
        "Le split est group-aware : les vues Top et Bottom potentiellement associées restent "
        "dans le même jeu. Sans cette contrainte, une vue d'une châtaigne pourrait servir à "
        "l'apprentissage et l'autre à l'évaluation, ce qui rendrait les scores artificiellement "
        "optimistes. Les contrôles existants confirment des intersections train/validation/test "
        "vides et l'intégrité de tous les groupes."
    )
    report.callout(
        "Protection du test",
        "le test n'a servi ni au choix de l'architecture, ni aux hyperparamètres, au learning "
        "rate, au checkpoint ou au seuil. Il a été ouvert une seule fois après gel complet de "
        "la configuration.",
        "success",
    )

    report.heading("4.2.3 Prétraitement et augmentation", level=1)
    report.paragraph(
        "Chaque image est lue en RGB, redimensionnée par interpolation bilinéaire en conservant "
        "son ratio, puis centrée dans un carré 224×224 par padding noir. Elle est ensuite "
        "convertie en tenseur et normalisée avec les statistiques ImageNet."
    )
    report.table(
        "Prétraitement commun et justification",
        ["Étape", "Application", "Justification CastagNet"],
        [
            ["RGB", "Tous les jeux", "Format d'entrée homogène à trois canaux."],
            ["224×224", "Tous les jeux", "Format standard torchvision, compromis information/coût."],
            ["Ratio conservé", "Tous les jeux", "Évite de déformer la forme du fruit."],
            ["Padding noir centré", "Tous les jeux", "Obtient un carré sans étirement artificiel."],
            ["Normalisation ImageNet", "Tous les jeux", "Cohérente avec les poids pré-entraînés."],
            ["Flip horizontal p=0,5", "Train uniquement", "Variation réaliste sans changer la classe."],
            ["Rotation ±10°", "Train uniquement", "Simule une faible variation d'orientation."],
            ["Aucune augmentation", "Validation et test", "Mesure déterministe et comparable."],
        ],
        "training/transforms.py et configurations JSON",
    )
    report.paragraph(
        "Les augmentations restent volontairement légères. Des rotations fortes ou des "
        "transformations photométriques non justifiées auraient pu créer des textures ou des "
        "positions absentes du dispositif réel. La validation et le test utilisent strictement "
        "le même pipeline déterministe, sans flip ni rotation."
    )

    report.heading("4.2.4 Baseline CNN maison", level=1, page_break=True)
    report.paragraph(
        "Un CNN maison sert de référence sans transfer learning. Son objectif n'est pas "
        "d'être optimal, mais de mesurer ce qu'une architecture légère entraînée depuis zéro "
        "peut apprendre sur CastagNet avant de mobiliser des représentations ImageNet."
    )
    report.paragraph(
        "L'architecture enchaîne quatre blocs Conv-BatchNorm-ReLU-MaxPool, avec 32, 64, 128 "
        "puis 256 canaux. Un AdaptiveAvgPool ramène les cartes à une valeur par canal, suivi "
        "d'un dropout de 0,3 et d'une couche linéaire 256→4. Aucun Softmax n'est placé dans "
        "forward, car CrossEntropyLoss attend directement les logits."
    )
    report.table(
        "Architecture du CNN maison",
        ["Composant", "Configuration", "Rôle"],
        [
            ["Bloc 1", "Conv 3→32, BN, ReLU, Pool", "Premiers motifs visuels"],
            ["Bloc 2", "Conv 32→64, BN, ReLU, Pool", "Motifs intermédiaires"],
            ["Bloc 3", "Conv 64→128, BN, ReLU, Pool", "Caractéristiques plus abstraites"],
            ["Bloc 4", "Conv 128→256, BN, ReLU, Pool", "Représentation haute dimension"],
            ["Sortie", "AdaptiveAvgPool, Dropout 0,3, Linear 256→4", "Classification multiclasse"],
        ],
        "training/models/custom_cnn.py",
    )
    report.paragraph(
        f"Le réseau compte {integer(smoke['CustomCNN v2']['model_parameters']['total_parameters'])} "
        f"paramètres, soit {smoke['CustomCNN v2']['model_parameters']['fp32_size_megabytes']:.4f} Mio "
        "en FP32. Cette petite taille en fait une baseline utile, mais limite aussi sa capacité "
        "à extraire les nuances visuelles entre classes proches."
    )
    report.paragraph(
        "CrossEntropyLoss non pondérée est adaptée aux quatre classes exclusives et garantit "
        "une comparaison identique entre expériences. AdamW a été choisi pour sa mise à jour "
        "robuste et sa séparation claire du weight decay. Le LR initial 0,001 constitue un "
        "point de départ standard pour cette baseline avec AdamW ; ReduceLROnPlateau et "
        "l'early stopping limitent les epochs inutiles."
    )

    report.heading("4.2.5 Analyse des deux versions du CNN", level=1)
    cnn_rows = []
    for name, lr in (("CustomCNN v1", "0,001"), ("CustomCNN v2", "0,0003")):
        values = metrics[name]
        val = values["best_validation"]
        cnn_rows.append([
            name, lr, str(values["best_epoch"]), pct(val["accuracy"]), pct(val["macro_f1"]),
            pct(val["per_class"]["Conforme"]["precision"]),
            pct(val["per_class"]["Conforme"]["recall"]),
            pct(val["per_class"]["NON Conforme"]["recall"]),
            pct(val["per_class"]["PIETRA"]["recall"]),
            str(val["business_errors"]["false_conforme_total"]),
        ])
    report.table(
        "Comparaison des deux CNN maison sur validation",
        ["Modèle", "LR", "Best ep.", "Accuracy", "Macro-F1", "P Conf.", "R Conf.", "R NC", "R PIETRA", "Faux Conf."],
        cnn_rows,
        "metrics_summary.json des runs CustomCNN v1 et v2",
    )
    report.figure(
        "Losses d'entraînement et de validation des deux CNN maison",
        figures["custom_loss"],
        "training_history.csv des runs CustomCNN v1 et v2",
    )
    report.paragraph(
        "La v1 atteint un rappel Conforme élevé, mais au prix d'une précision faible et de "
        "821 défauts NON Conforme ou PIETRA acceptés comme Conforme. Sa val_loss oscille "
        "fortement. L'audit du pipeline a validé model.eval(), l'absence de gradients en "
        "validation, shuffle=False, la remise à zéro des métriques, le calcul pondéré de la "
        "loss et l'intégrité du checkpoint : aucune fuite ou erreur d'évaluation n'a été trouvée."
    )
    report.paragraph(
        "Deux défauts de reporting ont néanmoins été corrigés pour les expériences suivantes : "
        "le LR était enregistré après scheduler.step, et le diagnostic d'overfitting comparait "
        "des epochs différentes. Ils n'ont modifié ni les poids ni les résultats v1."
    )
    report.paragraph(
        "La v2 ne change que le LR, abaissé à 0,0003 pour tester si une optimisation moins "
        "agressive stabilise l'apprentissage. Elle améliore la val_loss, l'accuracy, la "
        "macro-F1 et le rappel NON Conforme ; elle est donc retenue comme baseline maison. "
        "La v1 conserve légèrement moins de faux Conforme (821 contre 835), ce qui nuance "
        "le gain global de la v2."
    )
    report.callout(
        "Limite principale",
        "dans la v2, le rappel PIETRA n'est que de 11,60 % et 626 PIETRA sur 905 sont classées "
        "Conforme. Ce résultat indique d'abord une capacité de représentation insuffisante du "
        "petit CNN ; il ne démontre pas à lui seul un dataset défectueux.",
        "risk",
    )

    report.heading("4.2.6 Mise en place du transfer learning", level=1, page_break=True)
    report.paragraph(
        "Le dossier training annoncé dans le sujet était absent des données réellement "
        "fournies. Un pipeline simple et reproductible a donc été construit avec PyTorch, "
        "torchvision et MLflow, en conservant strictement les splits, la cible effective_label, "
        "le preprocessing, la seed 20260902 et les métriques."
    )
    report.paragraph(
        "Les poids ImageNet officiels torchvision fournissent des filtres visuels génériques "
        "déjà appris sur un vaste corpus. Pour un dataset d'environ 35 000 images, leur "
        "réutilisation permet d'éviter de réapprendre toutes les représentations depuis zéro "
        "et favorise une convergence plus efficace."
    )
    report.table(
        "Protocole commun de transfer learning",
        ["Phase", "Paramètres entraînés", "Epochs", "LR", "Régulation", "Justification"],
        [
            ["1 — tête", "Tête seule ; backbone gelé", "3", "0,001", "AdamW", "Adapter rapidement la sortie aux 4 classes."],
            ["2 — fine-tuning", "Réseau entièrement dégelé", "15 max.", "0,0001", "Plateau + early stop 5", "Ajuster progressivement les poids ImageNet."],
        ],
        "training/train.py et configurations des modèles pré-entraînés",
    )
    report.paragraph(
        "Le LR plus faible en fine-tuning évite de détruire brutalement les représentations "
        "pré-entraînées. CrossEntropyLoss reste non pondérée afin d'isoler l'effet de "
        "l'architecture et de conserver un protocole comparable. Aucun grand balayage "
        "d'hyperparamètres n'a été entrepris, ce qui limite le risque de sur-optimisation sur validation."
    )

    report.heading("4.2.7 Modèles pré-entraînés étudiés", level=1)
    architecture_rows = [
        ["CustomCNN v2", integer(389924), "1,4874 Mio", "Oui", "Baseline légère"],
        ["MobileNetV3 Small", integer(1521956), "5,8058 Mio", "Oui", "Représentant edge"],
        ["EfficientNet-B0", integer(4012672), "15,3071 Mio", "Oui", "Compromis intermédiaire"],
        ["ResNet18", integer(11178564), "42,6428 Mio", "Oui", "Candidat qualité"],
        ["ShuffleNetV2 x1.0", integer(1257704), "≈ 4,80 Mio", "Non", "Préparé uniquement"],
    ]
    report.table(
        "Complexité théorique et statut des architectures",
        ["Modèle", "Paramètres", "Taille FP32", "Entraîné", "Positionnement"],
        architecture_rows,
        "smoke_test_summary.json, factory torchvision et configurations",
    )
    report.paragraph(
        "MobileNetV3 Small représente l'option légère pertinente pour une future cible edge. "
        "EfficientNet-B0 fournit un point intermédiaire entre taille et qualité. ResNet18, plus "
        "lourd, sert de candidat orienté qualité. Ces rôles permettent une comparaison utile "
        "sans supposer qu'un modèle plus petit est automatiquement plus rapide sur le matériel réel."
    )
    report.paragraph(
        "ShuffleNetV2 x1.0 a été préparé mais volontairement non entraîné. MobileNet couvrait "
        "déjà le besoin d'un représentant très léger ; la comparaison comprenait alors une "
        "baseline, un modèle edge, un intermédiaire et un candidat qualité. Une expérience "
        "supplémentaire aurait apporté une valeur limitée au regard du temps et accru le risque "
        "de recherche excessive sur validation."
    )

    report.heading("4.2.8 Suivi des expériences avec MLflow", level=1)
    report.paragraph(
        "MLflow centralise configurations, hyperparamètres, learning rates, losses, accuracy, "
        "macro-F1, checkpoints, durées et artefacts. Il permet de comparer les modèles avec une "
        "trace reproductible et relie explicitement le checkpoint final au run qui l'a produit."
    )
    mlflow_rows = []
    for name in ("CustomCNN v1", "CustomCNN v2", "MobileNetV3 Small", "EfficientNet-B0", "ResNet18"):
        values = metrics[name]
        mlflow_rows.append([
            name,
            values["run_id"],
            str(values["epochs_completed"]),
            str(values["best_epoch"]),
            duration(values["training_time_seconds"]),
            mlflow[name]["status"],
        ])
    report.table(
        "Traçabilité des runs d'entraînement MLflow",
        ["Modèle", "Run ID", "Epochs", "Best ep.", "Durée", "Statut"],
        mlflow_rows,
        "mlflow.db et metrics_summary.json",
    )
    report.paragraph(
        "Pour ResNet18, l'historique distingue explicitement les trois epochs tête seule des "
        "quinze epochs de fine-tuning. Le checkpoint de l'epoch 13 minimise la val_loss à "
        f"{metrics['ResNet18']['best_val_loss']:.6f}. Le run final est "
        f"{metrics['ResNet18']['run_id']}."
    )
    report.figure(
        "ResNet18 — losses train et validation, phases et meilleur checkpoint",
        figures["resnet_loss"],
        "analysis/output/modeling/resnet18/training_history.csv",
    )
    report.figure(
        "ResNet18 — accuracy train et validation",
        figures["resnet_accuracy"],
        "analysis/output/modeling/resnet18/training_history.csv",
    )
    report.figure(
        "ResNet18 — macro-F1 de validation",
        figures["resnet_f1"],
        "analysis/output/modeling/resnet18/training_history.csv",
    )
    report.paragraph(
        "Les courbes montrent le gain du fine-tuning après adaptation de la tête. La hausse "
        "finale de la performance train alors que la validation plafonne justifie la sélection "
        "du meilleur checkpoint plutôt que de la dernière epoch. L'écart d'accuracy au meilleur "
        "checkpoint reste limité à 4,57 points selon l'artefact."
    )

    report.heading("4.2.9 Comparaison des architectures", level=1, page_break=True)
    comparison_rows = []
    for name in ("CustomCNN v2", "MobileNetV3 Small", "EfficientNet-B0", "ResNet18"):
        val = metrics[name]["best_validation"]
        comparison_rows.append([
            name,
            integer(smoke[name]["model_parameters"]["total_parameters"]),
            pct(val["accuracy"]),
            pct(val["macro_f1"]),
            pct(val["per_class"]["Conforme"]["precision"]),
            pct(val["per_class"]["Conforme"]["recall"]),
            pct(val["per_class"]["NON Conforme"]["recall"]),
            pct(val["per_class"]["PIETRA"]["recall"]),
            str(val["business_errors"]["false_conforme_total"]),
            str(val["business_errors"]["conforme_rejected"]),
        ])
    report.table(
        "Comparaison principale sur validation",
        ["Modèle", "Param.", "Acc.", "Macro-F1", "P Conf.", "R Conf.", "R NC", "R PIETRA", "Faux Conf.", "Conf. rejetés"],
        comparison_rows,
        "metrics_summary.json de chaque modèle",
    )
    report.paragraph(
        "Le transfer learning produit un saut qualitatif net. Par rapport au CustomCNN v2, "
        "MobileNet améliore déjà fortement la macro-F1 et la reconnaissance de PIETRA. "
        "EfficientNet progresse encore, puis ResNet18 obtient la meilleure accuracy, la meilleure "
        "macro-F1, la meilleure précision et le meilleur rappel Conforme, le meilleur rappel "
        "PIETRA et le plus faible compteur métier de faux Conforme."
    )
    report.paragraph(
        "MobileNetV3 Small reste attractif par sa taille, mais son rappel Conforme de "
        f"{pct(metrics['MobileNetV3 Small']['best_validation']['per_class']['Conforme']['recall'])} "
        "est inférieur au minimum de 85 %. EfficientNet-B0 constitue un compromis plus performant, "
        "avec un rappel Conforme de "
        f"{pct(metrics['EfficientNet-B0']['best_validation']['per_class']['Conforme']['recall'])}, "
        "mais demeure légèrement derrière ResNet18 sur les critères prioritaires."
    )

    report.heading("4.2.10 Sélection de ResNet18", level=1)
    report.paragraph(
        "ResNet18 est retenu comme meilleur candidat prédictif du §4.2. Il est le seul modèle "
        "pré-entraîné évalué dont le rappel Conforme dépasse 85 % avec la décision argmax, tout "
        "en obtenant les meilleurs indicateurs globaux et le plus faible nombre de défauts "
        "NON Conforme ou PIETRA acceptés comme Conforme."
    )
    report.paragraph(
        f"Le modèle compte {integer(smoke['ResNet18']['model_parameters']['total_parameters'])} "
        "paramètres, contre 1,52 million pour MobileNet et 4,01 millions pour EfficientNet. "
        "Cette différence n'est pas masquée : ResNet18 est choisi comme candidat qualité parce "
        "que la protection de la classe Conforme est prioritaire. Son coût embarqué devra "
        "néanmoins être démontré, et non supposé."
    )
    report.callout(
        "Portée de la sélection",
        "ResNet18 est retenu comme candidat qualité à l'issue du §4.2. Son adéquation définitive "
        "à la cible matérielle reste conditionnée aux benchmarks d'inférence du §4.3.",
        "warning",
    )

    report.heading("4.2.11 Calibration du seuil Conforme", level=1, page_break=True)
    report.paragraph(
        "La calibration a été conduite uniquement sur les 5 289 images de validation, sans "
        "réentraînement. Pour chaque seuil de 0,50 à 0,99 par pas de 0,01, une image est classée "
        "Conforme si P(Conforme) atteint le seuil ; sinon, l'argmax est calculé seulement parmi "
        "NON Conforme, PIETRA et Vide."
    )
    row055 = calibration.loc[np.isclose(calibration.threshold, .55)].iloc[0]
    row056 = calibration.loc[np.isclose(calibration.threshold, .56)].iloc[0]
    row091 = calibration.loc[np.isclose(calibration.threshold, .91)].iloc[0]
    report.table(
        "Trois seuils de décision examinés sur validation",
        ["Seuil", "P Conforme", "R Conforme", "Macro-F1", "Faux Conf.", "Conf. rejetés", "Décision"],
        [
            ["0,55", pct(row055.precision_conforme), pct(row055.recall_conforme), pct(row055.macro_f1), str(int(row055.false_conforme_total)), str(int(row055.conforme_rejected)), "Retenu"],
            ["0,56", pct(row056.precision_conforme), pct(row056.recall_conforme), pct(row056.macro_f1), str(int(row056.false_conforme_total)), str(int(row056.conforme_rejected)), "Rappel < 85 %"],
            ["0,91", pct(row091.precision_conforme), pct(row091.recall_conforme), pct(row091.macro_f1), str(int(row091.false_conforme_total)), str(int(row091.conforme_rejected)), "Rappel insuffisant"],
        ],
        "conforme_threshold_curve.csv — validation uniquement",
    )
    report.figure(
        "Calibration du seuil de décision Conforme",
        figures["calibration"],
        "analysis/output/modeling/resnet18/calibration/conforme_threshold_curve.csv",
    )
    report.paragraph(
        f"Aucun des {calibration_summary['thresholds_tested']} seuils ne satisfait simultanément "
        "95 % de précision et 85 % de rappel. Le compromis mathématique minimal se situe à 0,56, "
        f"avec {pct(row056.precision_conforme)} de précision et {pct(row056.recall_conforme)} "
        "de rappel. Comme le rappel passe sous le minimum métier, ce seuil est rejeté."
    )
    report.paragraph(
        f"Le seuil 0,55 conserve {pct(row055.recall_conforme)} de rappel pour "
        f"{pct(row055.precision_conforme)} de précision. Le gain de précision apporté par 0,56 "
        "est négligeable face au franchissement de la contrainte de rappel ; 0,55 a donc été "
        "figé avant l'ouverture du test."
    )
    report.paragraph(
        f"Le premier seuil atteignant 95 % de précision est 0,91, mais le rappel chute à "
        f"{pct(row091.recall_conforme)}. Ce compromis est incompatible avec l'exigence de "
        "rendement minimal. Plus généralement, le chevauchement des scores entre Conforme et "
        "les autres classes empêche un simple seuil d'atteindre les deux objectifs."
    )

    report.heading("4.2.12 Évaluation finale sur le jeu de test", level=1, page_break=True)
    report.paragraph(
        "Le test a été ouvert une seule fois après gel du modèle ResNet18, du checkpoint de "
        "l'epoch 13, du run MLflow et du seuil 0,55. Les 5 288 images ont fait l'objet d'une "
        "passe d'inférence unique ; les décisions argmax et seuil réutilisent exactement les "
        "mêmes probabilités. Aucun entraînement, recalibrage ou changement n'a suivi."
    )
    val = metrics["ResNet18"]["best_validation"]
    argmax = test["argmax"]
    threshold = test["threshold_055"]
    report.table(
        "ResNet18 argmax — validation et test",
        ["Jeu", "Accuracy", "Macro-F1", "P Conforme", "R Conforme", "Faux Conf.", "Conf. rejetés"],
        [
            ["Validation", pct(val["accuracy"]), pct(val["macro_f1"]), pct(val["per_class"]["Conforme"]["precision"]), pct(val["per_class"]["Conforme"]["recall"]), str(val["business_errors"]["false_conforme_total"]), str(val["business_errors"]["conforme_rejected"])],
            ["Test", pct(argmax["accuracy"]), pct(argmax["macro_f1"]), pct(argmax["per_class"]["Conforme"]["precision"]), pct(argmax["per_class"]["Conforme"]["recall"]), str(argmax["business_errors"]["false_conforme_total"]), str(argmax["business_errors"]["conforme_rejected"])],
        ],
        "ResNet18 metrics_summary.json et final_test_metrics.json",
    )
    report.paragraph(
        "La proximité des deux lignes indique une généralisation cohérente sur ce dataset : "
        "l'accuracy diminue d'environ 0,44 point et la macro-F1 de 0,68 point. Cette stabilité "
        "ne doit toutefois pas être extrapolée à de nouvelles campagnes, d'autres lots ou au "
        "fonctionnement en production."
    )
    report.table(
        "Test final — argmax et règle métier figée",
        ["Décision", "Accuracy", "Macro-P", "Macro-R", "Macro-F1", "P Conf.", "R Conf.", "Faux Conf.", "Conf. rejetés"],
        [
            ["Argmax", pct(argmax["accuracy"]), pct(argmax["macro_precision"]), pct(argmax["macro_recall"]), pct(argmax["macro_f1"]), pct(argmax["per_class"]["Conforme"]["precision"]), pct(argmax["per_class"]["Conforme"]["recall"]), str(argmax["business_errors"]["false_conforme_total"]), str(argmax["business_errors"]["conforme_rejected"])],
            ["Seuil 0,55", pct(threshold["accuracy"]), pct(threshold["macro_precision"]), pct(threshold["macro_recall"]), pct(threshold["macro_f1"]), pct(threshold["per_class"]["Conforme"]["precision"]), pct(threshold["per_class"]["Conforme"]["recall"]), str(threshold["business_errors"]["false_conforme_total"]), str(threshold["business_errors"]["conforme_rejected"])],
        ],
        "analysis/output/modeling/resnet18/final_test/final_test_metrics.json",
    )
    report.table(
        "Métriques par classe — configuration finale test",
        ["Classe", "Précision", "Rappel", "F1", "Support"],
        [
            [name, pct(values["precision"]), pct(values["recall"]), pct(values["f1"]), integer(values["support"])]
            for name, values in threshold["per_class"].items()
        ],
        "final_test_metrics.json — threshold_055",
    )
    report.figure(
        "Matrice de confusion ResNet18 sur test — argmax",
        figures["test_argmax"],
        "final_test_metrics.json — argmax",
        11.2,
    )
    report.figure(
        "Matrice de confusion ResNet18 sur test — seuil Conforme 0,55",
        figures["test_threshold"],
        "final_test_metrics.json — threshold_055",
        11.2,
    )
    report.paragraph(
        "Avec le seuil 0,55, 36 NON Conforme et 130 PIETRA sont acceptées comme Conforme, "
        "contre respectivement 42 et 146 en argmax. Le compteur métier baisse donc de 188 à "
        "166, au prix de 31 rejets Conforme supplémentaires. Ce déplacement est cohérent avec "
        "la priorité qualité : l'acceptation devient plus stricte sans dégrader fortement les "
        "métriques globales."
    )
    report.paragraph(
        "La confusion dominante reste Conforme↔PIETRA : 200 vrais Conforme sont classés PIETRA "
        "et 130 PIETRA deviennent Conforme avec la règle finale. La classe Vide reste très bien "
        "séparée. Le compteur « faux Conforme métier » additionne volontairement NON Conforme "
        "et PIETRA ; les trois Vide prédits Conforme restent inclus dans la précision statistique."
    )

    report.heading("4.2.13 Analyse du respect du cahier des charges", level=1)
    conf = threshold["per_class"]["Conforme"]
    report.table(
        "Respect des contraintes métier sur le test final",
        ["Critère", "Objectif", "Résultat", "Écart", "Verdict"],
        [
            ["Précision Conforme", "≥ 95 %", pct(conf["precision"]), f"−{(0.95-conf['precision'])*100:.2f} pts".replace(".", ","), "NON ATTEINT"],
            ["Rappel Conforme", "≥ 85 %", pct(conf["recall"]), f"+{(conf['recall']-0.85)*100:.2f} pts".replace(".", ","), "ATTEINT"],
            ["Respect simultané", "Deux critères", "Un critère sur deux", "—", "NON ATTEINT"],
        ],
        "final_test_metrics.json — business_requirements",
    )
    report.callout(
        "Conclusion obligatoire",
        "Le modèle final respecte l'objectif minimal de rappel de la classe Conforme mais ne "
        "satisfait pas l'objectif de précision de 95 %.",
        "risk",
    )
    report.paragraph(
        "Ce résultat est présenté sans ajustement post-test. Rechercher maintenant un autre "
        "seuil, checkpoint ou modèle à partir de ces 5 288 images transformerait le test en jeu "
        "de validation et invaliderait l'estimation finale. Le test est donc fermé à toute "
        "optimisation future."
    )

    report.heading("4.2.14 Prise en compte des contraintes d'embarquement", level=1)
    report.paragraph(
        "La complexité théorique place ResNet18 au-dessus de MobileNetV3 Small et EfficientNet-B0. "
        "Ce constat fournit une indication de stockage et de mémoire des poids, mais ne permet "
        "pas de conclure sur la latence réelle : les opérateurs, le runtime, la précision "
        "numérique et le matériel influencent directement le débit."
    )
    report.paragraph(
        "MobileNet est plus petit, mais n'est pas automatiquement plus rapide dans toutes les "
        "conditions. Inversement, les 42,64 Mio FP32 de ResNet18 ne prouvent pas qu'il est "
        "incompatible avec 3 Go de VRAM. Le §4.3 devra exporter en ONNX et mesurer taille, "
        "latence batch 1, débit, mémoire CPU/GPU et capacité à traiter douze flux sur ou au plus "
        "près de la GTX 1060 3 Go et de l'ancien i7."
    )

    report.heading("4.2.15 Limites, recul critique et pistes d'amélioration", level=1, page_break=True)
    report.heading("Limites observées", level=2)
    report.bullets([
        "La précision Conforme finale de 89,16 % reste inférieure aux 95 % demandés.",
        "PIETRA demeure la classe la plus difficile et partage une frontière visuelle avec Conforme.",
        "Les labels proviennent d'un processus humain historiquement hétérogène, malgré l'adjudication du §4.1.",
        "ResNet18 est le plus lourd des trois modèles pré-entraînés réellement évalués.",
        "La performance ONNX, la GTX 1060 3 Go et les douze flux n'ont pas encore été testés.",
        "Un seul dataset est disponible ; la robustesse sur une collecte indépendante reste inconnue.",
        "Le dossier training officiel annoncé dans le sujet était absent des données fournies.",
        "Le nombre d'architectures et le tuning ont été volontairement limités.",
        "Le test final est désormais fermé à toute optimisation.",
    ])
    report.paragraph(
        "Les ambiguïtés Conforme/PIETRA/NON Conforme peuvent provenir à la fois de la difficulté "
        "visuelle, de la variabilité des acquisitions et des conventions de labellisation. "
        "L'amélioration spectaculaire entre CNN maison et transfer learning montre toutefois "
        "que la capacité du modèle expliquait une part importante des erreurs initiales."
    )
    report.heading("Pistes futures — hors résultats du §4.2", level=2)
    report.bullets([
        "Harmoniser la labellisation et revoir de façon ciblée les PIETRA difficiles et les faux Conforme.",
        "Étudier une CrossEntropy pondérée, un coût métier asymétrique ou un WeightedRandomSampler.",
        "Évaluer une calibration probabiliste plus avancée sur de nouvelles données de validation.",
        "Collecter un jeu indépendant et étudier une approche hiérarchique Vide/Fruit puis qualité.",
        "Tester une autre résolution d'entrée uniquement dans un nouveau protocole expérimental.",
        "Évaluer quantification, compression et éventuellement distillation après le benchmark ONNX.",
        "Mesurer sur le matériel cible et sur douze flux avant tout choix embarqué définitif.",
    ])
    report.paragraph(
        "Les expériences s'arrêtent ici de manière raisonnée. Deux CNN maison, trois modèles "
        "pré-entraînés, trois niveaux de complexité, une calibration validation et un test final "
        "unique suffisent pour répondre à l'objectif méthodologique. Multiplier les essais "
        "apporterait une valeur décroissante et favoriserait la sur-optimisation sur validation."
    )

    report.heading("4.2.16 Conclusion et transition vers le §4.3", level=1)
    report.paragraph(
        "Le §4.2 a établi un pipeline reproductible malgré l'absence du dossier training annoncé, "
        "construit une baseline CNN, démontré l'intérêt du transfer learning et comparé des "
        "architectures de complexités différentes sous un protocole commun. MLflow relie chaque "
        "résultat à sa configuration et au checkpoint correspondant."
    )
    report.paragraph(
        "ResNet18 est le meilleur candidat prédictif sur validation. Le seuil 0,55, choisi avant "
        "le test, réduit les défauts acceptés comme Conforme tout en maintenant le rappel au-dessus "
        "de 85 %. L'évaluation finale reste cohérente avec la validation, mais la précision de "
        "95 % n'est pas atteinte : le cahier des charges est donc seulement partiellement satisfait."
    )
    report.paragraph(
        "La prochaine étape ne consiste pas à ajuster le modèle au test. Le §4.3 devra vérifier "
        "si le candidat qualité peut être exporté en ONNX et respecter les contraintes de "
        "latence, débit et mémoire sur la cible envisagée. Cette mesure décidera si ResNet18 "
        "peut être retenu tel quel ou si une optimisation embarquée, voire un modèle plus léger, "
        "doit être étudié sur un nouveau protocole."
    )

    report.heading("Annexe — Registre des artefacts utilisés", level=1, page_break=True)
    report.paragraph(
        "Les valeurs du rapport proviennent exclusivement des artefacts suivants. Les nombres "
        "arrondis dans le texte sont calculés à partir de leurs valeurs exactes."
    )
    report.bullets([
        "analysis/output/final_training_manifest_summary.json",
        "analysis/output/ml_split_summary.json",
        "analysis/output/modeling/*/config.json",
        "analysis/output/modeling/*/smoke_test_summary.json",
        "analysis/output/modeling/*/metrics_summary.json",
        "analysis/output/modeling/*/training_history.csv",
        "analysis/output/modeling/resnet18/calibration/conforme_threshold_curve.csv",
        "analysis/output/modeling/resnet18/calibration/conforme_threshold_summary.json",
        "analysis/output/modeling/resnet18/final_test/final_test_metrics.json",
        "analysis/output/modeling/resnet18/final_test/confusion_matrix_argmax.csv",
        "analysis/output/modeling/resnet18/final_test/confusion_matrix_threshold_055.csv",
        "mlflow.db — expérience CastagNet_4_2 et runs cités.",
    ])
    report.heading("Points de cohérence et conventions", level=2)
    report.bullets([
        "Les artefacts exacts priment sur les valeurs arrondies du cahier de consignes.",
        "Le compromis automatisé de calibration était 0,56 ; le choix métier final documenté est 0,55 afin de préserver le rappel minimal.",
        "Le diagnostic overfitting_probable de CustomCNN v1 provient d'un ancien défaut de reporting ; les poids et métriques ne sont pas affectés.",
        "Le compteur faux Conforme total des artefacts additionne NON Conforme et PIETRA ; les Vide→Conforme sont néanmoins comptés dans la précision.",
        "ShuffleNetV2 x1.0 est configuré mais ne possède aucun résultat d'entraînement.",
    ])
    report.save()
    metadata = {
        "docx": str(DOCX_PATH),
        "markdown": str(MD_PATH),
        "figures": [str(path) for path in figures.values()],
        "figure_titles": report.figure_titles,
        "table_titles": report.table_titles,
        "mlflow_runs": {name: values["run_id"] for name, values in mlflow.items()},
    }
    (OUTPUT / "report_4_2_generation_summary.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return report, metadata


def main() -> None:
    _, metadata = build_report()
    print(json.dumps(metadata, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
