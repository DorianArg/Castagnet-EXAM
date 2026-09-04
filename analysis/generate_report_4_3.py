"""Generate the CastagNet section 4.3 report from existing ONNX artifacts only."""

from __future__ import annotations

import json
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
MODELING = ROOT / "analysis" / "output" / "modeling"
RESNET = MODELING / "resnet18"
STATIC_ONNX = RESNET / "onnx"
BATCH_DIR = STATIC_ONNX / "batch_benchmark"
COMPARISON_DIR = MODELING / "onnx_comparison"
FIGURE_DIR = ROOT / "analysis" / "output" / "report_4_3_figures"
REPORTS = ROOT / "reports"
DOCX_PATH = REPORTS / "Rapport_CastagNet_4_3_ONNX_Embarquement_FINAL.docx"
MD_PATH = REPORTS / "Rapport_CastagNet_4_3_ONNX_Embarquement_FINAL.md"
SUMMARY_PATH = ROOT / "analysis" / "output" / "report_4_3_generation_summary.json"

NAVY = "17324D"
BLUE = "245B78"
TEAL = "2A7F7F"
GREEN = "4D8B6A"
AMBER = "C58520"
RED = "A94442"
PALE_BLUE = "EAF2F7"
PALE_TEAL = "E8F3F1"
PALE_AMBER = "FFF4DC"
PALE_RED = "FBEAEA"
DARK = "263238"
MID = "5D6A72"
WHITE = "FFFFFF"


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def fr(value: float, digits: int = 2) -> str:
    return f"{value:.{digits}f}".replace(".", ",")


def pct(value: float, digits: int = 2) -> str:
    return f"{value * 100:.{digits}f} %".replace(".", ",")


def integer(value: int) -> str:
    return f"{value:,}".replace(",", " ")


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_row_split(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def repeat_header(row) -> None:
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(node)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, text, separate, end))


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.75)
    section.right_margin = Cm(1.75)
    section.header_distance = Cm(0.75)
    section.footer_distance = Cm(0.75)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color, before, after in (
        ("Title", 28, NAVY, 0, 10),
        ("Subtitle", 15, TEAL, 0, 8),
        ("Heading 1", 17, NAVY, 14, 7),
        ("Heading 2", 13.5, BLUE, 11, 5),
        ("Heading 3", 11, TEAL, 8, 4),
    ):
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    if "Report Caption" not in [style.name for style in document.styles]:
        caption = document.styles.add_style("Report Caption", WD_STYLE_TYPE.PARAGRAPH)
        caption.font.name = "Aptos"
        caption.font.size = Pt(8.3)
        caption.font.italic = True
        caption.font.color.rgb = RGBColor.from_string(MID)
        caption.paragraph_format.space_before = Pt(2)
        caption.paragraph_format.space_after = Pt(7)
        caption.paragraph_format.keep_with_next = False
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run("CASTAGNET   |   §4.3 ONNX et embarquement")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Rapport final — 4 septembre 2026    •    ")
    add_field(footer, "PAGE")
    for paragraph in (header, footer):
        for run in paragraph.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string(MID)


class Report:
    def __init__(self) -> None:
        self.doc = Document()
        configure_document(self.doc)
        self.md: list[str] = []
        self.figure_number = 0
        self.table_number = 0
        self.figures: list[str] = []
        self.figure_titles: list[str] = []
        self.table_titles: list[str] = []

    def title_page(self) -> None:
        self.md += [
            "# CASTAGNET — §4.3 Export ONNX et évaluation embarquée", "",
            "**Rapport final — faisabilité, coût d'inférence et limites de déploiement**", "",
            "Mise en Situation Professionnelle — MSc BIHAR  ",
            "4 septembre 2026", "",
        ]
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(48)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("CASTAGNET")
        run.bold = True
        run.font.name = "Aptos Display"
        run.font.size = Pt(19)
        run.font.color.rgb = RGBColor.from_string(TEAL)
        paragraph = self.doc.add_paragraph(style="Title")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run("Export ONNX et\névaluation embarquée")
        paragraph = self.doc.add_paragraph(style="Subtitle")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run("Rapport final — §4.3")
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(28)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("Mise en Situation Professionnelle — MSc BIHAR\n4 septembre 2026")
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor.from_string(MID)
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(46)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(
            "Configuration figée\nResNet18 • checkpoint epoch 13 • ONNX opset 17\n"
            "Machine mesurée : RTX 3080 Laptop 8 Go — cible : GTX 1060 3 Go"
        )
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor.from_string(NAVY)
        self.doc.add_page_break()

    def heading(self, text: str, level: int = 1, page_break: bool = False) -> None:
        if page_break:
            self.doc.add_page_break()
        self.doc.add_heading(text, level=level)
        self.md += ["#" * (level + 1) + " " + text, ""]

    def paragraph(self, text: str) -> None:
        self.doc.add_paragraph(text)
        self.md += [text, ""]

    def bullets(self, items: list[str]) -> None:
        for item in items:
            paragraph = self.doc.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(2.5)
            paragraph.add_run(item)
            self.md.append(f"- {item}")
        self.md.append("")

    def callout(self, title: str, text: str, kind: str = "info") -> None:
        fill = {"info": PALE_BLUE, "success": PALE_TEAL, "warning": PALE_AMBER, "risk": PALE_RED}[kind]
        table = self.doc.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        set_cell_shading(cell, fill)
        set_cell_margins(cell, 125, 150, 125, 150)
        paragraph = cell.paragraphs[0]
        run = paragraph.add_run(title + " — ")
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(RED if kind == "risk" else NAVY)
        paragraph.add_run(text)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)
        self.md += [f"> **{title} —** {text}", ""]

    def table(self, title: str, headers: list[str], rows: list[list[str]], source: str, font_size: float = 8.4) -> None:
        self.table_number += 1
        full_title = f"Tableau {self.table_number} — {title}"
        self.table_titles.append(full_title)
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(full_title)
        run.bold = True
        run.font.size = Pt(9.3)
        run.font.color.rgb = RGBColor.from_string(NAVY)
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        header = table.rows[0]
        repeat_header(header)
        for index, value in enumerate(headers):
            cell = header.cells[index]
            cell.text = value
            set_cell_shading(cell, NAVY)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(WHITE)
                run.font.size = Pt(font_size)
        for row_index, values in enumerate(rows):
            row = table.add_row()
            prevent_row_split(row)
            for index, value in enumerate(values):
                cell = row.cells[index]
                cell.text = str(value)
                set_cell_margins(cell)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                if row_index % 2:
                    set_cell_shading(cell, "F5F8FA")
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        run.font.size = Pt(font_size)
        caption = self.doc.add_paragraph(style="Report Caption")
        caption.add_run(f"Source : {source}")
        self.md += [f"**{full_title}**", "", "| " + " | ".join(headers) + " |", "|" + "|".join(["---"] * len(headers)) + "|"]
        for values in rows:
            self.md.append("| " + " | ".join(str(value).replace("\n", "<br>") for value in values) + " |")
        self.md += ["", f"*Source : {source}*", ""]

    def figure(self, path: Path, title: str, source: str, width_cm: float = 15.7) -> None:
        self.figure_number += 1
        full_title = f"Figure {self.figure_number} — {title}"
        self.figure_titles.append(full_title)
        self.figures.append(str(path.resolve()))
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run().add_picture(str(path), width=Cm(width_cm))
        caption = self.doc.add_paragraph(style="Report Caption")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.add_run(f"{full_title}. Source : {source}")
        relative = path.relative_to(ROOT).as_posix()
        self.md += [f"![{full_title}]({relative})", "", f"*{full_title}. Source : {source}*", ""]

    def save(self) -> None:
        properties = self.doc.core_properties
        properties.title = "CastagNet — §4.3 Export ONNX et évaluation embarquée"
        properties.subject = "Faisabilité de déploiement de ResNet18 avec ONNX Runtime"
        properties.author = "Projet CastagNet"
        REPORTS.mkdir(parents=True, exist_ok=True)
        self.doc.save(DOCX_PATH)
        MD_PATH.write_text("\n".join(self.md), encoding="utf-8")


def plot_bar(labels, values, title, ylabel, path, color=TEAL, value_format="{:.2f}") -> None:
    fig, axis = plt.subplots(figsize=(8, 4.5))
    bars = axis.bar(labels, values, color="#" + color, width=0.58)
    axis.set_title(title, fontsize=13, fontweight="bold", color="#" + NAVY)
    axis.set_ylabel(ylabel)
    axis.grid(axis="y", alpha=0.22)
    axis.spines[["top", "right"]].set_visible(False)
    for bar, value in zip(bars, values):
        axis.text(bar.get_x() + bar.get_width() / 2, bar.get_height(), value_format.format(value), ha="center", va="bottom", fontsize=9)
    fig.tight_layout()
    fig.savefig(path, dpi=180, bbox_inches="tight")
    plt.close(fig)


def generate_figures(comparison: pd.DataFrame, batches: list[dict]) -> list[Path]:
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    labels = ["ResNet18", "EfficientNet-B0", "MobileNetV3 Small"]
    order = ["resnet18", "efficientnet_b0", "mobilenet_v3_small"]
    frame = comparison.set_index("model").loc[order]
    figures = []
    specifications = [
        ("cpu_inference_latency.png", frame.cpu_inference_mean_ms.tolist(), "Latence d'inférence ONNX sur CPU", "Latence moyenne (ms)", TEAL),
        ("cuda_inference_latency.png", frame.cuda_inference_mean_ms.tolist(), "Latence d'inférence ONNX sur CUDA", "Latence moyenne (ms)", BLUE),
        ("cuda_inference_throughput.png", frame.cuda_inference_images_per_second.tolist(), "Débit d'inférence ONNX sur CUDA", "Images/seconde", GREEN),
        ("onnx_model_size.png", frame.onnx_size_mib.tolist(), "Taille des modèles ONNX", "Mio", AMBER),
    ]
    for filename, values, title, ylabel, color in specifications:
        path = FIGURE_DIR / filename
        plot_bar(labels, values, title, ylabel, path, color)
        figures.append(path)

    path = FIGURE_DIR / "quality_vs_cuda_latency.png"
    fig, axis = plt.subplots(figsize=(8, 4.8))
    axis.scatter(frame.cuda_inference_mean_ms, frame.macro_f1 * 100, s=np.sqrt(frame.parameters) * 0.22, c=["#245B78", "#2A7F7F", "#C58520"], alpha=0.85)
    for label, x, y in zip(labels, frame.cuda_inference_mean_ms, frame.macro_f1 * 100):
        axis.annotate(label, (x, y), xytext=(6, 6), textcoords="offset points", fontsize=9)
    axis.set_title("Compromis macro-F1 / latence CUDA", fontsize=13, fontweight="bold", color="#" + NAVY)
    axis.set_xlabel("Latence CUDA moyenne (ms) — plus faible est préférable")
    axis.set_ylabel("Macro-F1 validation (%) — plus élevée est préférable")
    axis.grid(alpha=0.22)
    axis.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(path, dpi=180, bbox_inches="tight")
    plt.close(fig)
    figures.append(path)

    batch_sizes = [row["batch_size"] for row in batches]
    batch_specs = [
        ("batch_inference_throughput.png", [row["inference_only"]["images_per_second"] for row in batches], "Débit d'inférence ResNet18 selon le batch", "Images/seconde", GREEN),
        ("batch_latency_per_image.png", [row["inference_only"]["mean_ms_per_image"] for row in batches], "Latence d'inférence ramenée à une image", "ms/image", BLUE),
        ("batch_e2e_throughput.png", [row["preprocessing_plus_inference"]["images_per_second"] for row in batches], "Débit end-to-end ResNet18 selon le batch", "Images/seconde", TEAL),
    ]
    for filename, values, title, ylabel, color in batch_specs:
        path = FIGURE_DIR / filename
        fig, axis = plt.subplots(figsize=(8, 4.5))
        axis.plot(batch_sizes, values, marker="o", linewidth=2.2, color="#" + color)
        if filename == "batch_e2e_throughput.png":
            axis.axhline(60, color="#" + RED, linestyle="--", linewidth=1.3, label="12 caméras × 5 FPS = 60 img/s")
            axis.legend(fontsize=8)
        for x, value in zip(batch_sizes, values):
            axis.annotate(fr(value, 2), (x, value), xytext=(0, 7), textcoords="offset points", ha="center", fontsize=8.5)
        axis.set_title(title, fontsize=13, fontweight="bold", color="#" + NAVY)
        axis.set_xlabel("Taille du batch")
        axis.set_ylabel(ylabel)
        axis.set_xticks(batch_sizes)
        axis.grid(alpha=0.22)
        axis.spines[["top", "right"]].set_visible(False)
        fig.tight_layout()
        fig.savefig(path, dpi=180, bbox_inches="tight")
        plt.close(fig)
        figures.append(path)

    batch12 = next(row for row in batches if row["batch_size"] == 12)
    path = FIGURE_DIR / "batch12_stage_times.png"
    stage_labels = ["Décodage", "Preprocessing", "Constitution\ndu batch", "session.run"]
    values = [batch12[key]["mean_ms_per_batch"] for key in ("decode_only", "preprocessing_without_decode", "batch_construction", "inference_only")]
    plot_bar(stage_labels, values, "Temps moyens isolés au batch 12", "ms/batch", path, BLUE)
    figures.append(path)
    return figures


def validate_artifacts() -> tuple[dict, list[str]]:
    paths = {
        "export ResNet18 statique": STATIC_ONNX / "export_summary.json",
        "validation détaillée statique": STATIC_ONNX / "numerical_validation_details.csv",
        "smoke benchmark ResNet18": STATIC_ONNX / "benchmark_smoke.json",
        "benchmark ResNet18 statique": STATIC_ONNX / "benchmark_results.json",
        "comparaison ONNX JSON": COMPARISON_DIR / "onnx_comparison.json",
        "comparaison ONNX CSV": COMPARISON_DIR / "onnx_comparison.csv",
        "rapport comparatif ONNX": COMPARISON_DIR / "onnx_comparison_report.md",
        "export ResNet18 dynamique": BATCH_DIR / "dynamic_export_summary.json",
        "benchmark batch JSON": BATCH_DIR / "batch_benchmark.json",
        "benchmark batch CSV": BATCH_DIR / "batch_benchmark.csv",
        "rapport benchmark batch": BATCH_DIR / "batch_benchmark_report.md",
        "résultats finaux §4.2": RESNET / "final_test" / "final_test_metrics.json",
        "modèle ONNX statique": STATIC_ONNX / "resnet18_castagnet.onnx",
        "modèle ONNX dynamique": STATIC_ONNX / "resnet18_castagnet_dynamic_batch.onnx",
        "script export": ROOT / "training" / "export_onnx.py",
        "script benchmark statique": ROOT / "training" / "benchmark_onnx.py",
        "script benchmark batch": ROOT / "training" / "benchmark_dynamic_batch.py",
        "script comparaison": ROOT / "training" / "compare_onnx.py",
    }
    missing = [str(path) for path in paths.values() if not path.is_file()]
    if missing:
        raise FileNotFoundError(f"Artefacts manquants: {missing}")
    static_export = load_json(paths["export ResNet18 statique"])
    static_benchmark = load_json(paths["benchmark ResNet18 statique"])
    comparison_json = load_json(paths["comparaison ONNX JSON"])
    dynamic_export = load_json(paths["export ResNet18 dynamique"])
    batch_benchmark = load_json(paths["benchmark batch JSON"])
    final_test = load_json(paths["résultats finaux §4.2"])
    checks = {
        "static_export_valid": static_export["status"] == "valid",
        "dynamic_export_valid": dynamic_export["status"] == "valid",
        "static_test_not_loaded": static_benchmark["test_manifest_loaded"] is False,
        "comparison_test_not_loaded": comparison_json["test_manifest_loaded"] is False,
        "batch_test_not_loaded": batch_benchmark["test_manifest_loaded"] is False,
        "batch_training_not_performed": batch_benchmark["training_performed"] is False,
        "checkpoint_epoch_13": static_export["checkpoint_epoch"] == dynamic_export["checkpoint_epoch"] == 13,
        "static_sha256_matches": static_export["onnx_sha256"] == "e663bf4b6f3746072ea978f53ba6d958f3ea49c591c5a97311bebea219bcfd23",
        "static_size_matches": static_export["onnx_size_bytes"] == 44704751,
        "all_static_providers_succeeded": static_benchmark["provider_failures"] == {},
        "batch_row_count": len(batch_benchmark["batch_results"]) == 5,
        "test_threshold_frozen": final_test["frozen_conforme_threshold"] == 0.55,
    }
    if not all(checks.values()):
        raise RuntimeError(f"Contrôles des artefacts échoués: {checks}")
    return {
        "paths": paths,
        "static_export": static_export,
        "static_benchmark": static_benchmark,
        "comparison_json": comparison_json,
        "comparison": pd.read_csv(paths["comparaison ONNX CSV"]),
        "dynamic_export": dynamic_export,
        "batch_benchmark": batch_benchmark,
        "final_test": final_test,
        "checks": checks,
    }, [
        "Les mesures batch 1 statique et batch 1 dynamique proviennent de campagnes distinctes; leurs faibles écarts ne doivent pas être fusionnés.",
        "Les mesures isolées des étapes et la mesure end-to-end reposent sur des boucles différentes; elles ne sont pas strictement additives.",
        "La VRAM propre au processus est absente des artefacts et ne peut pas être chiffrée.",
    ]


def build_report(data: dict, inconsistencies: list[str], figures: list[Path]) -> Report:
    report = Report()
    static_export = data["static_export"]
    static_benchmark = data["static_benchmark"]
    comparison = data["comparison"].set_index("model")
    dynamic_export = data["dynamic_export"]
    batches = data["batch_benchmark"]["batch_results"]
    final_test = data["final_test"]
    cpu = static_benchmark["results"]["cpu"]
    cuda = static_benchmark["results"]["cuda"]
    threshold_test = final_test["threshold_055"]

    report.title_page()
    report.heading("Synthèse décisionnelle", 1)
    report.callout(
        "Réponse à la question du §4.3",
        "ResNet18 est raisonnablement exécutable avec ONNX Runtime sur la machine de développement. Son inférence CUDA est rapide, mais la faisabilité sur la GTX 1060 3 Go et les 12 flux réels n'est pas démontrée. Le goulot observé est principalement la préparation CPU des images.",
        "success",
    )
    report.paragraph(
        "Le choix prédictif du §4.2 reste cohérent pour un déploiement GPU : ResNet18 combine la meilleure macro-F1 de validation (85,21 %) et la meilleure latence CUDA batch 1 mesurée parmi les trois architectures. En revanche, la chaîne complète atteint 52,47 images/s en batch 1 et au mieux 58,04 images/s avec batching, soit moins que les 60 images/s requises par douze caméras à 5 FPS."
    )
    report.bullets([
        "Mesure réelle : Windows 11, CPU Intel64 Family 6 Model 165, 16 CPU logiques, RTX 3080 Laptop 8 Go.",
        "Cible non mesurée : GTX 1060 3 Go, ancien Intel i7 et douze flux caméra.",
        "Résultat prédictif figé du §4.2 : précision Conforme 89,16 %, rappel Conforme 85,54 %, macro-F1 84,44 % sur test avec seuil 0,55.",
        "Priorité future : optimiser et paralléliser le décodage et le preprocessing avant de chercher uniquement un réseau plus petit.",
    ])

    report.heading("4.3.1 Objectifs et contraintes matérielles", 1, page_break=True)
    report.paragraph(
        "Le §4.3 évalue si le modèle retenu au §4.2 peut être intégré dans une chaîne d'inférence embarquée. Il ne cherche ni à modifier les poids ni à améliorer les métriques de classification. L'étude porte sur l'export, la fidélité numérique, les temps CPU/GPU, le coût end-to-end et le potentiel de mutualisation pour douze caméras."
    )
    report.table(
        "Configuration figée héritée du §4.2",
        ["Élément", "Valeur", "Justification"],
        [
            ["Modèle", "ResNet18", "Meilleur compromis prédictif mesuré au §4.2"],
            ["Checkpoint", "Epoch 13", "Meilleure loss de validation du run retenu"],
            ["Run MLflow", "db0ee7873903411e899e2f85384ae58a", "Traçabilité des poids"],
            ["Entrée / sortie", "RGB 224 × 224 / 4 logits", "Conforme au preprocessing et au mapping entraînés"],
            ["Mapping", "Conforme 0 ; NON Conforme 1 ; PIETRA 2 ; Vide 3", "Ordre figé des classes"],
            ["Seuil métier", "P(Conforme) = 0,55, hors graphe", "Règle figée sur validation et indépendante du réseau"],
        ],
        "config.json, export_summary.json et final_test_metrics.json",
    )
    report.callout(
        "Cahier des charges prédictif",
        f"Sur le test final déjà ouvert au §4.2, le rappel Conforme atteint {pct(threshold_test['per_class']['Conforme']['recall'])}, mais la précision de {pct(threshold_test['per_class']['Conforme']['precision'])} reste sous l'objectif de 95 %. Le §4.3 conserve ce résultat sans ajustement.",
        "warning",
    )
    machine = static_benchmark["machine_measured"]
    report.table(
        "Machine mesurée et matériel cible",
        ["Statut", "CPU / système", "GPU", "Portée"],
        [
            ["Mesure réelle", f"Windows 11 ; {machine['logical_cpu_count']} CPU logiques\nIntel64 Family 6 Model 165", "RTX 3080 Laptop\n8 192 Mio ; pilote 591.74", "Tous les chiffres de latence du rapport"],
            ["Cible du sujet", "Ancien Intel i7", "GTX 1060 3 Go", "Analyse qualitative uniquement ; aucune extrapolation"],
        ],
        "benchmark_results.json",
    )

    report.heading("4.3.2 Export ONNX de ResNet18", 1)
    report.paragraph(
        "ONNX découple le réseau du framework PyTorch utilisé pendant l'entraînement et permet son exécution avec ONNX Runtime sur CPU ou GPU. L'opset 17 couvre les opérateurs de ResNet18 et est compatible avec la version du runtime utilisée ; des opérateurs plus récents n'étaient pas nécessaires."
    )
    report.table(
        "Caractéristiques de l'export ONNX principal",
        ["Propriété", "Valeur", "Contrôle"],
        [
            ["Fichier", "resnet18_castagnet.onnx", "Export batch fixe conservé"],
            ["Entrée", "input : [1, 3, 224, 224]", "Batch 1, RGB"],
            ["Sortie", "logits : [1, 4]", "Quatre classes"],
            ["Opset", str(static_export["opset"]), "Validé par ONNX checker"],
            ["Taille", f"{integer(static_export['onnx_size_bytes'])} octets\n{fr(static_export['onnx_size_mib'])} Mio", "Poids FP32 exportés"],
            ["Softmax", "Absent", "Le graphe expose les logits"],
            ["Seuil 0,55", "Absent", "Post-traitement métier externe"],
            ["SHA-256", static_export["onnx_sha256"], "Intégrité de l'artefact"],
        ],
        "export_summary.json et modèle ONNX",
        8.0,
    )
    report.paragraph(
        "Le Softmax et le seuil ne sont pas intégrés : le réseau reste un classifieur générique produisant des logits, tandis que la règle métier demeure explicite et modifiable indépendamment du modèle. Cette séparation évite de confondre calcul neuronal et décision opérationnelle."
    )

    report.heading("4.3.3 Validation numérique PyTorch / ONNX", 1)
    numerical = static_export["numerical_validation"]
    report.paragraph(
        "La fidélité a été contrôlée sur 16 images déterministes, régulièrement espacées dans le manifeste de validation. Ce choix rend le contrôle reproductible sans rouvrir le test final. Les sorties ne sont pas bit-à-bit identiques, mais les écarts FP32 sont de très faible amplitude."
    )
    report.table(
        "Fidélité de l'export ResNet18 batch 1",
        ["Indicateur", "Valeur", "Interprétation"],
        [
            ["Max. écart absolu logits", f"{numerical['max_absolute_difference_logits']:.7e}", "Inférieur à la tolérance"],
            ["Écart absolu moyen logits", f"{numerical['mean_absolute_difference_logits']:.7e}", "Très faible"],
            ["Max. écart probabilités", f"{numerical['max_absolute_difference_probabilities']:.7e}", "Très faible"],
            ["Écart moyen probabilités", f"{numerical['mean_absolute_difference_probabilities']:.7e}", "Très faible"],
            ["Argmax différents", str(numerical["argmax_prediction_differences"]), "Aucune classe modifiée"],
            ["Décisions seuil 0,55 différentes", str(numerical["threshold_055_decision_differences"]), "Aucune décision métier modifiée"],
        ],
        "export_summary.json et numerical_validation_details.csv",
    )
    report.callout("Verdict", "L'export ONNX est fidèle à PyTorch sur l'échantillon contrôlé et peut être utilisé pour les mesures d'inférence.", "success")

    report.heading("4.3.4 Protocole de benchmark", 1, page_break=True)
    report.table(
        "Choix méthodologiques et justification",
        ["Choix", "Application", "Pourquoi dans CastagNet ?"],
        [
            ["Batch 1", "Référence embarquée", "Représente la décision immédiate sur une image"],
            ["50 warmups", "Avant chaque série", "Écarte les coûts d'initialisation non représentatifs du régime stabilisé"],
            ["500 itérations", "Benchmark statique", "Réduit la sensibilité aux variations ponctuelles"],
            ["Moyenne, médiane, p95, p99", "Distribution des latences", "Les percentiles montrent les ralentissements que la moyenne masque"],
            ["Inference-only / end-to-end", "Mesures séparées", "Distingue le coût du réseau du coût applicatif réel"],
            ["CPU / CUDA", "Deux providers", "Compare repli CPU et accélération GPU sur la machine disponible"],
            ["TensorRT non retenu", "Provider détecté seulement", "Aucun benchmark final TensorRT n'a été produit"],
        ],
        "benchmark_onnx.py et benchmark_results.json",
    )
    providers = ", ".join(static_benchmark["available_providers"])
    report.paragraph(
        f"Les providers détectés sont {providers}. Les références utilisent CPUExecutionProvider et CUDAExecutionProvider, qui fournissent une comparaison directe et reproductible. Le temps de création de session est exclu ; le décodage est inclus dans le preprocessing."
    )

    report.heading("4.3.5 Performances CPU et CUDA de ResNet18", 1)
    benchmark_rows = []
    for provider_name, result in (("CPU", cpu), ("CUDA", cuda)):
        for scope_name, key in (("Inférence seule", "inference_only"), ("End-to-end", "preprocessing_plus_inference")):
            values = result[key]
            benchmark_rows.append([
                provider_name, scope_name, fr(values["mean_ms"], 3), fr(values["median_ms"], 3),
                fr(values["p90_ms"], 3), fr(values["p95_ms"], 3), fr(values["p99_ms"], 3),
                fr(values["images_per_second"], 2),
            ])
    report.table(
        "Benchmark ResNet18 ONNX batch 1",
        ["Provider", "Périmètre", "Moy. ms", "Méd. ms", "p90", "p95", "p99", "img/s"],
        benchmark_rows,
        "benchmark_results.json — 50 warmups, 500 itérations",
        7.7,
    )
    report.paragraph(
        f"Sur CUDA, l'inférence seule prend {fr(cuda['inference_only']['mean_ms'], 3)} ms, contre {fr(cpu['inference_only']['mean_ms'], 3)} ms sur CPU. Le pipeline CUDA complet atteint toutefois {fr(cuda['preprocessing_plus_inference']['mean_ms'], 3)} ms : le preprocessing seul représente déjà {fr(cuda['preprocessing']['mean_ms'], 3)} ms en moyenne. La RAM maximale observée est de {fr(cpu['memory']['process_ram_peak_observed_mib'])} Mio sur CPU et {fr(cuda['memory']['process_ram_peak_observed_mib'])} Mio pendant la série CUDA ; la VRAM du processus n'a pas été obtenue."
    )
    report.figure(figures[0], "Latence d'inférence CPU des trois modèles", "onnx_comparison.csv")
    report.figure(figures[1], "Latence d'inférence CUDA des trois modèles", "onnx_comparison.csv")
    report.paragraph(
        "Les mesures séparées ne doivent pas être additionnées pour reconstruire l'end-to-end : les boucles, allocations et états de cache diffèrent. La mesure combinée peut inclure décodage, conversions, constitution des tenseurs, copies mémoire, orchestration Python et synchronisations CUDA. Faute d'instrumentation plus fine dans ce benchmark, aucune cause unique n'est attribuée. L'end-to-end reste la référence applicative."
    )

    report.heading("4.3.6 Comparaison ResNet18 / EfficientNet / MobileNet", 1, page_break=True)
    report.paragraph(
        "Les trois architectures effectivement entraînées au §4.2 ont été exportées avec le même format afin de vérifier si un réseau plus léger procure un avantage réel. La taille et le nombre de paramètres ne suffisent pas à anticiper la latence d'un runtime donné."
    )
    comparison_rows = []
    display_names = {"resnet18": "ResNet18", "efficientnet_b0": "EfficientNet-B0", "mobilenet_v3_small": "MobileNetV3 Small"}
    for model in ("resnet18", "efficientnet_b0", "mobilenet_v3_small"):
        row = comparison.loc[model]
        comparison_rows.append([
            display_names[model], pct(row.macro_f1), integer(int(row.parameters)), fr(row.onnx_size_mib),
            f"{fr(row.cpu_inference_mean_ms, 3)} ms\n{fr(row.cpu_inference_images_per_second, 2)} img/s",
            f"{fr(row.cuda_inference_mean_ms, 3)} ms\n{fr(row.cuda_inference_images_per_second, 2)} img/s",
            f"{fr(row.cpu_end_to_end_mean_ms, 3)} ms", f"{fr(row.cuda_end_to_end_mean_ms, 3)} ms",
        ])
    report.table(
        "Qualité de validation et coût ONNX",
        ["Modèle", "Macro-F1", "Paramètres", "ONNX Mio", "CPU inf.", "CUDA inf.", "CPU E2E", "CUDA E2E"],
        comparison_rows,
        "onnx_comparison.csv ; métriques validation figées du §4.2",
        7.5,
    )
    report.figure(figures[2], "Débit d'inférence CUDA des trois architectures", "onnx_comparison.csv")
    report.figure(figures[3], "Taille des trois modèles ONNX", "export_summary.json des trois architectures")

    report.heading("4.3.7 Analyse du compromis qualité / coût", 1)
    report.figure(figures[4], "Macro-F1 de validation et latence CUDA batch 1", "onnx_comparison.csv")
    report.callout(
        "Résultat CUDA",
        "ResNet18 est à la fois le meilleur prédictivement, le plus rapide en inférence seule et légèrement le plus rapide end-to-end sur la RTX 3080 mesurée. Un modèle plus petit n'est donc pas automatiquement plus rapide : l'efficacité dépend aussi des opérateurs, de leur parallélisation et du runtime.",
        "success",
    )
    report.paragraph(
        "Sur CPU, MobileNetV3 Small est nettement plus attractif : 4,834 ms par image contre 17,756 ms pour ResNet18, avec un fichier de 5,81 Mio. Ce gain s'accompagne d'une macro-F1 inférieure de 4,72 points. Le choix dépend donc du contexte : ResNet18 reste cohérent pour le GPU mesuré, tandis que MobileNet constitue une option crédible si le déploiement devient principalement CPU ou très contraint en stockage."
    )
    report.table(
        "Fidélité ONNX des modèles alternatifs",
        ["Modèle", "Max. écart logits", "Écart moyen logits", "Max. écart probabilités", "Argmax différents", "Verdict"],
        [
            ["EfficientNet-B0", "0,000931740", "0,000149926", "0,000073820", "0", "Valide"],
            ["MobileNetV3 Small", "0,000210762", "0,000039052", "0,000024498", "0", "Valide"],
        ],
        "export_summary.json de chaque architecture",
    )

    report.heading("4.3.8 Étude du batch dynamique", 1, page_break=True)
    report.paragraph(
        "Avec douze caméras, plusieurs images peuvent arriver dans une fenêtre proche. Un second export autorise donc N images par appel afin d'étudier le potentiel de mutualisation GPU. Il conserve exactement le checkpoint epoch 13, l'opset 17, les quatre logits et le post-traitement externe ; il ne remplace pas l'export batch 1."
    )
    report.table(
        "Validation numérique de l'export dynamique",
        ["Batch", "Max. écart logits", "Max. écart probabilités", "Argmax différents", "Décisions 0,55 différentes"],
        [[
            str(row["batch_size"]), f"{row['max_absolute_difference_logits']:.9f}",
            f"{row['max_absolute_difference_probabilities']:.9f}", str(row["argmax_differences"]),
            str(row["threshold_055_decision_differences"]),
        ] for row in dynamic_export["numerical_validation"]],
        "dynamic_export_summary.json — validation uniquement",
    )
    batch_rows = []
    for row in batches:
        inference = row["inference_only"]
        end_to_end = row["preprocessing_plus_inference"]
        batch_rows.append([
            str(row["batch_size"]), fr(inference["mean_ms_per_batch"], 3), fr(inference["p95_ms_per_batch"], 3),
            fr(inference["mean_ms_per_image"], 3), fr(inference["images_per_second"], 2),
            fr(end_to_end["mean_ms_per_batch"], 3), fr(end_to_end["mean_ms_per_image"], 3), fr(end_to_end["images_per_second"], 2),
        ])
    report.table(
        "Benchmark CUDA du batch dynamique",
        ["Batch", "Inf. ms/batch", "Inf. p95", "Inf. ms/image", "Inf. img/s", "E2E ms/batch", "E2E ms/image", "E2E img/s"],
        batch_rows,
        "batch_benchmark.json — 50 warmups, 300 itérations par taille",
        7.5,
    )
    report.figure(figures[5], "Débit d'inférence ResNet18 selon la taille du batch", "batch_benchmark.json")
    report.figure(figures[6], "Latence d'inférence ramenée à une image", "batch_benchmark.json")
    report.paragraph(
        "Entre les batches 1 et 12, le débit d'inférence passe de 243,02 à 1 527,94 images/s, soit environ ×6,3, et la latence par image de 4,115 à 0,654 ms. Le batch 12 est relié aux douze caméras pour étudier ce potentiel ; il ne représente pas une capture concurrente de douze flux."
    )

    report.heading("4.3.9 Analyse de la capacité 12 flux", 1)
    report.figure(figures[7], "Débit end-to-end selon le batch et besoin à 5 FPS/caméra", "batch_benchmark.json")
    report.table(
        "Besoins théoriques associés aux douze caméras",
        ["Cadence/caméra", "Besoin total", "Meilleur E2E mesuré", "Couverture sur machine mesurée"],
        [
            ["1 FPS", "12 img/s", "58,04 img/s", "Couverte théoriquement"],
            ["2 FPS", "24 img/s", "58,04 img/s", "Couverte théoriquement"],
            ["5 FPS", "60 img/s", "58,04 img/s", "Proche mais non couverte"],
            ["10 FPS", "120 img/s", "58,04 img/s", "Non couverte"],
            ["25 FPS", "300 img/s", "58,04 img/s", "Non couverte"],
        ],
        "batch_benchmark.json — calcul séquentiel simplifié",
    )
    report.callout(
        "Lecture correcte",
        "Le scénario 5 FPS par caméra est proche de la capacité mesurée mais reste légèrement supérieur au meilleur débit end-to-end observé : 60 img/s requis contre 58,04 img/s mesurés. Il n'est donc pas atteint.",
        "warning",
    )
    report.paragraph(
        "Le contraste est déterminant : l'inférence seule atteint 1 527,94 images/s au batch 12, alors que le meilleur pipeline complet atteint 58,04 images/s au batch 8. La performance du réseau seul ne représente donc pas la performance applicative."
    )

    report.heading("4.3.10 Identification des goulots d'étranglement", 1, page_break=True)
    batch12 = next(row for row in batches if row["batch_size"] == 12)
    stage_rows = []
    for label, key, scope in (
        ("Décodage", "decode_only", "Ouverture PIL, décodage, conversion RGB"),
        ("Preprocessing hors décodage", "preprocessing_without_decode", "Resize, padding, tenseur et normalisation"),
        ("Constitution du batch", "batch_construction", "Empilement NumPy"),
        ("session.run", "inference_only", "Copie hôte→GPU incluse, CUDA, synchronisation et sortie"),
    ):
        stage_rows.append([label, fr(batch12[key]["mean_ms_per_batch"], 3), scope])
    report.table("Mesures isolées au batch 12", ["Étape", "Moyenne ms/batch", "Périmètre"], stage_rows, "batch_benchmark.json")
    report.figure(figures[8], "Répartition des temps isolés au batch 12", "batch_benchmark.json")
    report.paragraph(
        "L'inférence ONNX Runtime représente moins de 8 ms pour douze images, tandis que décodage et preprocessing dépassent 140 ms dans les mesures isolées. Le principal goulot observé est donc la préparation CPU. Les valeurs isolées ne sont pas additives à l'end-to-end, car elles proviennent de boucles distinctes avec des allocations et états de cache différents."
    )
    report.callout(
        "Conséquence",
        "Remplacer ResNet18 par un modèle plus petit ne garantit pas une amélioration importante du pipeline complet. L'optimisation du décodage et du preprocessing est prioritaire.",
        "info",
    )

    report.heading("4.3.11 Discussion de la cible GTX 1060 3 Go", 1)
    report.paragraph(
        "Le fichier ResNet18 ONNX occupe seulement 42,63 Mio, ce qui indique que les poids seuls sont modestes devant 3 Go. Cette taille ne suffit toutefois pas à valider la mémoire d'exécution : activations, buffers, runtime CUDA, pipeline d'images et autres processus consomment également de la VRAM."
    )
    report.table(
        "Ce qui est validé et ce qui reste à valider sur la cible",
        ["Dimension", "État", "Conclusion autorisée"],
        [
            ["Fichier ONNX", "42,63 Mio mesurés", "Poids stockables ; mémoire runtime inconnue"],
            ["VRAM", "Non mesurée par processus", "Impossible d'affirmer que 3 Go suffisent"],
            ["Latence GPU", "RTX 3080 Laptop mesurée", "Aucune latence GTX 1060 déductible"],
            ["CPU", "CPU actuel mesuré", "Ancien i7 non validé"],
            ["Douze flux", "Calcul et batching synthétique", "Pas de validation réelle concurrente"],
            ["Système", "Windows 11", "Résultats dépendants de l'environnement mesuré"],
        ],
        "benchmark_results.json et limites expérimentales",
    )
    report.callout(
        "Limite essentielle",
        "Les mesures démontrent la faisabilité sur la machine de développement et identifient les goulots du pipeline, mais elles ne constituent pas une validation directe de la GTX 1060 3 Go.",
        "risk",
    )

    report.heading("4.3.12 Limites et recul critique", 1, page_break=True)
    report.bullets([
        "Le matériel cible n'était pas disponible ; la RTX 3080 Laptop est beaucoup plus récente que la GTX 1060.",
        "La VRAM propre au processus n'a pas été mesurée de manière fiable.",
        "Les entrées sont des fichiers image ; la capture et le décodage vidéo réels de douze flux ne sont pas reproduits intégralement.",
        "La concurrence, les files d'attente, l'ordonnancement et la contention CPU/GPU d'une application complète ne sont pas simulés.",
        "Le batching dynamique mesure une mutualisation potentielle, pas un système temps réel à douze caméras.",
        "Les résultats sont liés à Windows, aux versions d'ONNX Runtime et aux providers installés.",
        "Aucune extrapolation numérique n'est réalisée de la RTX 3080 vers la GTX 1060.",
    ])
    report.paragraph("Nuances relevées lors de la consolidation :")
    report.bullets(inconsistencies)
    report.heading("Points forts du protocole", 2)
    report.bullets([
        "Exports ONNX contrôlés par ONNX checker puis comparés numériquement à PyTorch.",
        "Mesures CPU et CUDA après warm-up, sur plusieurs centaines d'itérations avec moyenne, médiane et percentiles.",
        "Séparation entre inférence seule et chaîne end-to-end.",
        "Comparaison de trois architectures réellement entraînées et exportées.",
        "Validation du batch dynamique pour N = 1, 4 et 12, puis benchmark jusqu'au batch 12.",
        "Séparation explicite entre mesures réelles et analyse qualitative de la cible.",
    ])

    report.heading("4.3.13 Pistes d'optimisation", 1)
    report.paragraph("Les propositions suivantes sont des travaux futurs ; aucune n'a été mise en œuvre dans le cadre du §4.3.")
    report.table(
        "Priorités d'optimisation futures",
        ["Priorité", "Piste", "Motivation"],
        [
            ["1", "Optimiser décodage et preprocessing", "Goulot dominant observé"],
            ["2", "Paralléliser le preprocessing CPU", "Alimenter le GPU sans attente"],
            ["3", "Pipeline asynchrone et files d'attente", "Découpler capture, préparation et inférence"],
            ["4", "Prétraitement GPU si pertinent", "Réduire la charge CPU et les conversions"],
            ["5", "Batching dynamique piloté par délai", "Mutualiser sans dégrader excessivement la latence"],
            ["6", "Benchmark direct GTX 1060 et mesure VRAM", "Seule validation fiable de la cible"],
            ["7", "TensorRT, FP16 puis quantification", "À mesurer sur la cible, sans gain présumé"],
            ["8", "Échantillonnage adaptatif des frames", "Éviter plusieurs inférences sur la même châtaigne"],
        ],
        "Interprétation des benchmarks §4.3",
    )

    report.heading("4.3.14 Conclusion du §4.3", 1)
    report.paragraph(
        "ResNet18 est rapide en inférence CUDA sur la RTX 3080 Laptop mesurée et son export ONNX est fidèle au modèle PyTorch. La comparaison ne révèle aucune raison immédiate de le remplacer par EfficientNet-B0 ou MobileNetV3 Small dans ce contexte GPU ; MobileNet demeure cependant nettement plus intéressant sur CPU."
    )
    report.paragraph(
        "La conclusion principale porte sur le système complet : le batching améliore fortement l'utilisation du GPU, mais le débit end-to-end reste limité par le décodage et le preprocessing CPU. Les scénarios de douze caméras à 1 ou 2 FPS sont couverts théoriquement sur la machine actuelle ; 5 FPS par caméra reste légèrement hors de portée et les flux réels n'ont pas été simulés."
    )
    report.callout(
        "Arrêt raisonné des expériences",
        "L'étude dispose désormais d'un export vérifié, de mesures CPU/CUDA et end-to-end, d'une comparaison de trois architectures, d'un batch dynamique jusqu'à 12 et d'une identification du goulot. Sans le matériel cible, poursuivre apporterait peu de valeur et encouragerait des extrapolations fragiles. La prochaine expérience informative est un benchmark direct sur la GTX 1060 / i7 cible.",
        "success",
    )
    report.heading("Artefacts de référence", 2)
    report.bullets([str(path.relative_to(ROOT)) for path in data["paths"].values()])
    return report


def main() -> None:
    data, inconsistencies = validate_artifacts()
    figures = generate_figures(data["comparison"], data["batch_benchmark"]["batch_results"])
    report = build_report(data, inconsistencies, figures)
    report.save()
    summary = {
        "status": "completed",
        "docx": str(DOCX_PATH.resolve()),
        "markdown": str(MD_PATH.resolve()),
        "figures": report.figures,
        "figure_titles": report.figure_titles,
        "table_titles": report.table_titles,
        "artifacts_used": [str(path.resolve()) for path in data["paths"].values()],
        "checks": data["checks"],
        "inconsistencies_or_methodological_nuances": inconsistencies,
        "training_performed": False,
        "ml_evaluation_performed": False,
        "benchmark_performed": False,
    }
    SUMMARY_PATH.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "status": "completed",
        "docx": str(DOCX_PATH.resolve()),
        "markdown": str(MD_PATH.resolve()),
        "figures": len(report.figures),
        "tables": len(report.table_titles),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
